from __future__ import annotations

import hashlib
import base64
import io
import json
import os
import random
import re
import secrets
from difflib import SequenceMatcher
from math import sqrt
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote, quote_plus

from .course_catalog import ROLE_COURSES, build_seed_catalog
from .resume_parser import extract_profile

import requests
from requests import exceptions as request_errors
from requests.adapters import HTTPAdapter
from fastapi import Depends, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer
from jose import JWTError, jwt
from pydantic import BaseModel, EmailStr, Field, field_validator
from urllib3.util.retry import Retry

try:
    import certifi
except Exception:  # pragma: no cover
    certifi = None

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None

try:
    from pymongo import MongoClient
except Exception:  # pragma: no cover
    MongoClient = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

try:
    import fitz
    import pytesseract
    from PIL import Image
except Exception:  # pragma: no cover
    fitz = pytesseract = Image = None


APP_DIR = Path(__file__).resolve().parent
DATA_FILE = APP_DIR / "local_store.json"


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_env_file(APP_DIR / ".env")
JWT_SECRET = os.getenv("JWT_SECRET", "career-copilot-local-dev-secret")
JWT_ALGORITHM = "HS256"
ACCESS_TOKEN_MINUTES = 60 * 24
MONGO_URI = os.getenv("MONGO_URI", "")
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")
GITHUB_TOKEN = os.getenv("GITHUB_TOKEN", "")
GROQ_MODEL = os.getenv("GROQ_MODEL", "llama-3.1-8b-instant")
GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
VERIFY_SSL = os.getenv("GROQ_VERIFY_SSL", "true").lower() not in {"0", "false", "no"}
REQUEST_VERIFY = (certifi.where() if certifi else True) if VERIFY_SSL else False
GITHUB_TIMEOUT = (5, 30)
GROQ_PLACEHOLDERS = {"", "replace-with-your-groq-key", "your-groq-api-key"}
GROQ_AUTH_FAILED = False
MAX_RESUME_CHARS = 28000
MAX_PROFILE_CHARS = 52000

app = FastAPI(title="AI Career Copilot API", version="2.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://127.0.0.1:5173", "http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/api/auth/login")


class RegisterRequest(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    email: EmailStr
    password: str = Field(min_length=6, max_length=128)


class LoginRequest(BaseModel):
    email: EmailStr
    password: str


class ProfileRequest(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    target_role: str = Field(min_length=2, max_length=80)
    skills: list[str] = Field(default_factory=list, max_length=80)
    github_username: str = Field(default="", max_length=80)
    career_interests: str = Field(default="", max_length=800)

    @field_validator("skills")
    @classmethod
    def clean_skills(cls, value: list[str]) -> list[str]:
        return clean_list(value, 80)


class SkillRequest(BaseModel):
    skills: list[str] = Field(default_factory=list, max_length=80)
    career_interests: str = Field(default="", max_length=800)
    target_roles: list[str] = Field(default_factory=list, max_length=8)

    @field_validator("skills", "target_roles")
    @classmethod
    def clean_items(cls, value: list[str]) -> list[str]:
        return clean_list(value, 80)


class SkillGapRequest(BaseModel):
    current_skills: list[str] = Field(default_factory=list, max_length=80)
    target_career: str = Field(min_length=2, max_length=100)
    resume_text: str = Field(default="", max_length=MAX_RESUME_CHARS)

    @field_validator("current_skills")
    @classmethod
    def clean_items(cls, value: list[str]) -> list[str]:
        return clean_list(value, 80)


class InterviewStartRequest(BaseModel):
    round: str = Field(min_length=2, max_length=40)
    language: str = Field(default="Python", max_length=40)
    target_role: str = Field(default="", max_length=80)
    experience_level: str = Field(default="Fresher", max_length=40)


class InterviewAnswerRequest(BaseModel):
    session_id: str = Field(min_length=4, max_length=80)
    question_id: str = Field(min_length=4, max_length=120)
    answer: str = Field(min_length=3, max_length=8000)
    round: str = Field(min_length=2, max_length=40)
    language: str = Field(default="Python", max_length=40)
    code: str = Field(default="", max_length=16000)
    action: str = Field(default="submit", pattern="^(submit|skip|timeout|finish)$")
    time_taken: int = Field(default=0, ge=0, le=7200)


INTERVIEW_ROUNDS: tuple[tuple[str, str], ...] = (
    ("Introduction & Resume", "Easy"),
    ("Core Technical Fundamentals", "Easy"),
    ("Programming / Coding", "Medium"),
    ("Scenario Based Questions", "Medium"),
    ("Projects & Resume Discussion", "Medium"),
    ("Problem Solving", "Hard"),
    ("Advanced Concepts", "Hard"),
    ("Behavioral / HR", "Medium"),
    ("Final Company-style Rapid Fire", "Hard"),
)

INTERVIEW_COMPANIES = ("Google", "Microsoft", "Amazon", "Meta", "Netflix", "Adobe", "Zoho", "Infosys", "TCS", "Wipro", "Accenture", "Capgemini", "Cognizant")
QUESTIONS_PER_ROUND = 10
TOTAL_INTERVIEW_QUESTIONS = len(INTERVIEW_ROUNDS) * QUESTIONS_PER_ROUND


class GitHubAnalysisRequest(BaseModel):
    identifier: str = Field(min_length=1, max_length=200)
    target_role: str = Field(default="", max_length=80)


class MentorRequest(BaseModel):
    message: str = Field(min_length=2, max_length=4000)


class CourseProgressRequest(BaseModel):
    course_id: str = Field(min_length=2, max_length=160)
    progress: int = Field(default=0, ge=0, le=100)


def clean_list(items: list[str], limit: int) -> list[str]:
    seen: set[str] = set()
    cleaned: list[str] = []
    for item in items:
        value = re.sub(r"\s+", " ", str(item)).strip()
        key = value.lower()
        if value and key not in seen:
            cleaned.append(value[:limit])
            seen.add(key)
    return cleaned


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def hash_password(password: str) -> str:
    salt = secrets.token_hex(12)
    digest = hashlib.sha256(f"{salt}:{password}".encode()).hexdigest()
    return f"{salt}:{digest}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, digest = stored.split(":", 1)
    except ValueError:
        return False
    return hashlib.sha256(f"{salt}:{password}".encode()).hexdigest() == digest


def create_token(email: str) -> str:
    expires = datetime.now(timezone.utc) + timedelta(minutes=ACCESS_TOKEN_MINUTES)
    return jwt.encode({"sub": email, "exp": expires}, JWT_SECRET, algorithm=JWT_ALGORITHM)


class Store:
    def __init__(self) -> None:
        self.client = None
        self.db = None
        mongo_error = None
        is_vercel = os.getenv("VERCEL") == "1"
        if is_vercel and (not JWT_SECRET or JWT_SECRET == "career-copilot-local-dev-secret"):
            raise RuntimeError("Set JWT_SECRET to a unique random value for Vercel deployments.")
        if MONGO_URI and MongoClient:
            try:
                self.client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=1500)
                self.client.admin.command("ping")
                db_name = MONGO_URI.rsplit("/", 1)[-1].split("?", 1)[0]
                self.db = self.client[db_name or "career_copilot"]
            except Exception as exc:
                self.client = None
                self.db = None

                mongo_error = exc
        if is_vercel and self.db is None:
            message = "Set MONGO_URI to a reachable MongoDB database for Vercel deployments."
            if mongo_error:
                raise RuntimeError(message) from mongo_error
            raise RuntimeError(message)
        if self.db is None and not DATA_FILE.exists():
            DATA_FILE.write_text(json.dumps({"users": [], "courses": []}, indent=2), encoding="utf-8")
        self.seed_courses()

    @property
    def using_mongo(self) -> bool:
        return self.db is not None

    def _load(self) -> dict[str, Any]:
        data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
        data.setdefault("users", [])
        data.setdefault("courses", [])
        return data

    def _save(self, data: dict[str, Any]) -> None:
        DATA_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")

    def get_user(self, email: str) -> dict[str, Any] | None:
        email = email.lower()
        if self.db is not None:
            return self.db.users.find_one({"email": email}, {"_id": 0})
        return next((user for user in self._load()["users"] if user["email"] == email), None)

    def upsert_user(self, user: dict[str, Any]) -> None:
        user["email"] = user["email"].lower()
        if self.db is not None:
            self.db.users.update_one({"email": user["email"]}, {"$set": user}, upsert=True)
            return
        data = self._load()
        users = [item for item in data["users"] if item["email"] != user["email"]]
        users.append(user)
        data["users"] = users
        self._save(data)

    def seed_courses(self) -> None:
        catalog = build_seed_catalog()
        if self.db is not None:
            for course in catalog:
                self.db.course_catalog.update_one({"id": course["id"]}, {"$set": course}, upsert=True)
            self.db.course_catalog.create_index([("role", 1), ("roadmapOrder", 1)])
            return
        data = self._load()
        if data["courses"] != catalog:
            data["courses"] = catalog
            self._save(data)

    def courses_for_role(self, role: str) -> list[dict[str, Any]]:
        if self.db is not None:
            return list(self.db.course_catalog.find({"role": role}, {"_id": 0}).sort("roadmapOrder", 1))
        return [course for course in self._load()["courses"] if course.get("role") == role]


store = Store()


def default_user(name: str, email: str, password: str) -> dict[str, Any]:
    return {
        "name": name,
        "email": email.lower(),
        "password_hash": hash_password(password),
        "target_role": "Frontend Developer",
        "career_interests": "",
        "skills": ["HTML", "CSS", "JavaScript", "React"],
        "github_username": "",
        "resume_score": 0,
        "resume_text": "",
        "resume_sections": {},
        "mentor_messages": [],
        "interview_sessions": [],
        "course_progress": {},
        "certificates": [],
        "ai_cache": {},
        "created_at": now_iso(),
    }


def seed_demo_user() -> None:
    if not store.get_user("demo@careercopilot.ai"):
        store.upsert_user(default_user("Rahul Kumar", "demo@careercopilot.ai", "Demo@1234"))


seed_demo_user()


def current_user(token: str = Depends(oauth2_scheme)) -> dict[str, Any]:
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        email = payload.get("sub")
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid authentication token")
    user = store.get_user(email)
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    return user


def public_user(user: dict[str, Any]) -> dict[str, Any]:
    clean = dict(user)
    for key in ("password_hash", "ai_cache", "mentor_messages"):
        clean.pop(key, None)
    clean["database"] = "mongodb" if store.using_mongo else "local-json"
    return clean


def groq_configured() -> bool:
    return bool(GROQ_API_KEY and GROQ_API_KEY.strip() not in GROQ_PLACEHOLDERS)


def ensure_groq() -> None:
    if not groq_configured():
        raise HTTPException(status_code=503, detail="GROQ_API_KEY is required for AI features.")


def cache_key(name: str, payload: dict[str, Any]) -> str:
    raw = json.dumps(payload, sort_keys=True, ensure_ascii=False)
    return f"{name}:{hashlib.sha256(raw.encode()).hexdigest()}"


def cached(user: dict[str, Any], key: str) -> Any | None:
    item = user.get("ai_cache", {}).get(key)
    return item.get("data") if isinstance(item, dict) else None


def set_cached(user: dict[str, Any], key: str, data: Any) -> None:
    user.setdefault("ai_cache", {})[key] = {"created_at": now_iso(), "data": data}


def parse_json_content(content: str) -> dict[str, Any]:
    try:
        value = json.loads(content)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", content, flags=re.S)
        if not match:
            raise ValueError("Groq did not return valid JSON")
        value = json.loads(match.group(0))
    if not isinstance(value, dict):
        raise ValueError("Groq returned a non-object JSON value")
    return value


def groq_request_error(exc: requests.RequestException) -> HTTPException:
    message = str(exc)
    lowered = message.lower()
    if isinstance(exc, request_errors.Timeout):
        detail = "Groq request timed out. Check your internet connection and try again."
    elif isinstance(exc, request_errors.SSLError):
        detail = "Could not verify Groq's SSL certificate. Fix the local certificate store, or set GROQ_VERIFY_SSL=false only for local development."
    elif "failed to resolve" in lowered or "nameresolutionerror" in lowered or "getaddrinfo failed" in lowered:
        detail = "Could not resolve api.groq.com. Check DNS/internet access, VPN/proxy settings, or try restarting the backend after the network is back."
    elif isinstance(exc, request_errors.ConnectionError):
        detail = "Could not connect to Groq. Check internet access, firewall, VPN/proxy settings, and retry."
    else:
        detail = f"Groq request failed: {message}"
    return HTTPException(status_code=502, detail=detail)


def groq_error_detail(response: requests.Response) -> str:
    try:
        detail = response.json().get("error", {}).get("message", response.text)
    except ValueError:
        detail = response.text or response.reason
    if "invalid api key" in str(detail).lower():
        return "AI provider authentication failed. Update apps/api/.env and restart the backend."
    return str(detail)


def is_groq_auth_error(exc: HTTPException) -> bool:
    detail = str(exc.detail).lower()
    return exc.status_code in {502, 503} and (
        "ai provider authentication failed" in detail
        or "ai provider unavailable" in detail
        or "invalid groq api key" in detail
        or "groq_api_key is required" in detail
        or "invalid api key" in detail
    )


def groq_completion(
    messages: list[dict[str, str]],
    temperature: float = 0.25,
    response_format: dict[str, str] | None = None,
    max_tokens: int = 1800,
) -> str:
    global GROQ_AUTH_FAILED
    ensure_groq()
    if GROQ_AUTH_FAILED:
        raise HTTPException(status_code=503, detail="AI provider unavailable until backend restart.")
    payload: dict[str, Any] = {
        "model": GROQ_MODEL,
        "messages": messages,
        "temperature": temperature,
        "max_tokens": max_tokens,
    }
    if response_format:
        payload["response_format"] = response_format
    try:
        response = requests.post(
            GROQ_URL,
            headers={"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"},
            json=payload,
            timeout=35,
            verify=REQUEST_VERIFY,
        )
    except requests.RequestException as exc:
        raise groq_request_error(exc) from exc
    if not response.ok:
        detail = groq_error_detail(response)
        if "authentication failed" in detail.lower():
            GROQ_AUTH_FAILED = True
        raise HTTPException(status_code=502, detail=detail)
    try:
        content = response.json()["choices"][0]["message"]["content"]
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Invalid Groq response shape: {exc}") from exc
    if not isinstance(content, str) or not content.strip():
        raise HTTPException(status_code=502, detail="Groq returned an empty response.")
    return content.strip()


def groq_json(system: str, user_prompt: str, schema_hint: dict[str, Any], temperature: float = 0.25) -> dict[str, Any]:
    content = groq_completion(
        [
            {"role": "system", "content": system},
            {"role": "user", "content": user_prompt},
            {"role": "user", "content": f"Return only valid JSON matching this shape. Do not include markdown or commentary: {json.dumps(schema_hint)}"},
        ],
        temperature=temperature,
        response_format={"type": "json_object"},
        max_tokens=2400,
    )
    try:
        return parse_json_content(content)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Invalid Groq JSON response: {exc}") from exc


def profile_context(user: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": user.get("name", ""),
        "target_role": user.get("target_role", ""),
        "career_interests": user.get("career_interests", ""),
        "skills": user.get("skills", []),
        "education": user.get("resume_sections", {}).get("education", []),
        "experience": user.get("resume_sections", {}).get("experience", []),
        "projects": user.get("resume_sections", {}).get("projects", []),
        "resume_text": user.get("resume_text", "")[:MAX_PROFILE_CHARS],
        "github_username": user.get("github_username", ""),
        "github_analysis": user.get("github_analysis", {}),
    }


def mentor_context(user: dict[str, Any]) -> dict[str, Any]:
    """Return the saved career evidence the mentor is allowed to use."""
    interviews = user.get("mock_interviews", []) or user.get("interview_sessions", [])
    latest_interview = interviews[-1] if interviews else {}
    courses = recommended_courses(user)
    return {
        **profile_context(user),
        "resume_ats_report": {
            "score": int(user.get("resume_score") or 0),
            "sections": user.get("resume_sections", {}),
            "analysis": user.get("resume_analysis", {}),
        },
        "skill_gap_analysis": user.get("latest_skill_gap", {}),
        "career_readiness": int(
            (user.get("latest_skill_gap", {}) or {}).get("job_readiness_percentage") or 0
        ),
        "roadmap": (user.get("latest_skill_gap", {}) or {}).get("learning_roadmap", []),
        "latest_mock_interview": {
            key: latest_interview.get(key)
            for key in (
                "role", "status", "overall", "answers", "round_summaries",
                "strengths", "weaknesses", "areas_to_improve",
            )
            if key in latest_interview
        },
        "learning_progress": user.get("course_progress", {}),
        "recommended_courses": courses.get("courses", [])[:10],
        "certificates": user.get("certificates", []),
        "job_recommendations": user.get("job_recommendations", []),
    }


ROLE_SKILLS: dict[str, list[str]] = {
    "frontend": ["HTML", "CSS", "JavaScript", "TypeScript", "React", "State Management", "API Integration", "Responsive Design", "Testing", "Git"],
    "backend": ["Python", "FastAPI", "SQL", "REST APIs", "Authentication", "Data Modeling", "Testing", "Docker", "Cloud Deployment", "Git"],
    "full stack": ["HTML", "CSS", "JavaScript", "React", "Python", "FastAPI", "SQL", "Authentication", "Deployment", "Git"],
    "data": ["Python", "SQL", "Excel", "Pandas", "Statistics", "Data Visualization", "Dashboarding", "Business Analysis", "Git"],
    "ai": ["Python", "Machine Learning", "Pandas", "NumPy", "Scikit-learn", "Prompt Engineering", "Vector Databases", "APIs", "Model Evaluation", "Git"],
    "ml": ["Python", "Machine Learning", "Pandas", "NumPy", "Scikit-learn", "Deep Learning", "MLOps", "Model Evaluation", "Git"],
}

SKILL_OFFICIAL_RESOURCES: dict[str, str] = {
    "python": "https://docs.python.org/3/",
    "tensorflow": "https://www.tensorflow.org/tutorials",
    "fastapi": "https://fastapi.tiangolo.com/",
    "docker": "https://docs.docker.com/get-started/",
    "kubernetes": "https://kubernetes.io/docs/tutorials/",
    "sql": "https://www.w3schools.com/sql/",
    "git": "https://git-scm.com/doc",
    "langchain": "https://python.langchain.com/docs/introduction/",
    "html": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Structuring_content",
    "css": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Styling_basics",
    "javascript": "https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide",
    "typescript": "https://www.typescriptlang.org/docs/handbook/intro.html",
    "react": "https://react.dev/learn",
    "state management": "https://redux.js.org/tutorials/essentials/part-1-overview-concepts",
    "api integration": "https://developer.mozilla.org/en-US/docs/Web/API/Fetch_API/Using_Fetch",
    "apis": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Extensions/Client-side_APIs/Introduction",
    "rest apis": "https://developer.mozilla.org/en-US/docs/Glossary/REST",
    "responsive design": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/CSS_layout/Responsive_Design",
    "testing": "https://web.dev/learn/testing/",
    "authentication": "https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html",
    "data modeling": "https://learn.microsoft.com/en-us/power-bi/guidance/star-schema",
    "cloud deployment": "https://aws.amazon.com/getting-started/hands-on/",
    "deployment": "https://docs.docker.com/get-started/",
    "excel": "https://support.microsoft.com/en-us/excel",
    "pandas": "https://pandas.pydata.org/docs/getting_started/intro_tutorials/",
    "statistics": "https://developers.google.com/machine-learning/crash-course/prereqs-and-prework",
    "data visualization": "https://matplotlib.org/stable/tutorials/index.html",
    "dashboarding": "https://learn.microsoft.com/en-us/training/powerplatform/power-bi",
    "business analysis": "https://learn.microsoft.com/en-us/training/powerplatform/power-bi",
    "machine learning": "https://developers.google.com/machine-learning/crash-course",
    "numpy": "https://numpy.org/learn/",
    "scikit-learn": "https://scikit-learn.org/stable/tutorial/index.html",
    "prompt engineering": "https://www.promptingguide.ai/",
    "vector databases": "https://docs.pinecone.io/guides/get-started/overview",
    "model evaluation": "https://developers.google.com/machine-learning/crash-course/classification/accuracy-precision-recall",
    "deep learning": "https://www.tensorflow.org/tutorials/quickstart/beginner",
    "mlops": "https://cloud.google.com/architecture/mlops-continuous-delivery-and-automation-pipelines-in-machine-learning",
    "problem solving": "https://www.freecodecamp.org/learn/project-euler/",
    "communication": "https://learn.microsoft.com/en-us/training/modules/communicate-effectively/",
    "projects": "https://www.freecodecamp.org/learn/",
}

SKILL_LEARNING_META: dict[str, tuple[str, str]] = {
    "python": ("Beginner", "20 Hours"),
    "fastapi": ("Intermediate", "8 Hours"),
    "tensorflow": ("Advanced", "24 Hours"),
    "docker": ("Intermediate", "10 Hours"),
    "kubernetes": ("Advanced", "24 Hours"),
    "sql": ("Beginner", "16 Hours"),
    "git": ("Beginner", "8 Hours"),
    "langchain": ("Intermediate", "12 Hours"),
}


ROADMAP_RESOURCE_OVERRIDES: dict[str, dict[str, str]] = {
    "react": {
        "course": "https://www.freecodecamp.org/learn/front-end-development-libraries/",
        "practice": "https://www.frontendmentor.io/",
        "cheat_sheet": "https://react.dev/reference/react",
    },
    "fastapi": {
        "course": "https://www.freecodecamp.org/news/fastapi-tutorial/",
        "practice": "https://github.com/fastapi/full-stack-fastapi-template",
        "cheat_sheet": "https://fastapi.tiangolo.com/reference/",
    },
    "docker": {
        "practice": "https://labs.play-with-docker.com/",
        "cheat_sheet": "https://docs.docker.com/reference/cli/docker/",
    },
    "node.js": {
        "official": "https://nodejs.org/en/docs",
        "practice": "https://roadmap.sh/nodejs",
        "cheat_sheet": "https://nodejs.org/api/synopsis.html",
    },
    "nodejs": {
        "official": "https://nodejs.org/en/docs",
        "practice": "https://roadmap.sh/nodejs",
        "cheat_sheet": "https://nodejs.org/api/synopsis.html",
    },
    "express.js": {
        "official": "https://expressjs.com/",
        "practice": "https://github.com/expressjs/express",
        "cheat_sheet": "https://expressjs.com/en/4x/api.html",
    },
    "express": {
        "official": "https://expressjs.com/",
        "practice": "https://github.com/expressjs/express",
        "cheat_sheet": "https://expressjs.com/en/4x/api.html",
    },
    "mongodb": {
        "official": "https://www.mongodb.com/docs/",
        "course": "https://learn.mongodb.com/",
        "practice": "https://www.mongodb.com/developer/",
        "cheat_sheet": "https://www.mongodb.com/developer/products/mongodb/cheat-sheet/",
    },
    "sql": {
        "practice": "https://sqlbolt.com/",
        "cheat_sheet": "https://www.w3schools.com/sql/sql_quickref.asp",
    },
    "git": {
        "official": "https://git-scm.com/docs",
        "practice": "https://learngitbranching.js.org/",
        "cheat_sheet": "https://training.github.com/downloads/github-git-cheat-sheet/",
    },
    "javascript": {
        "official": "https://developer.mozilla.org/en-US/docs/Web/JavaScript",
        "practice": "https://javascript.info/",
        "cheat_sheet": "https://developer.mozilla.org/en-US/docs/Web/JavaScript/Reference",
    },
    "typescript": {
        "official": "https://www.typescriptlang.org/docs/",
        "practice": "https://www.totaltypescript.com/",
        "cheat_sheet": "https://www.typescriptlang.org/cheatsheets/",
    },
}


SKILL_ALIASES: dict[str, str] = {
    "node": "Node.js",
    "nodejs": "Node.js",
    "node.js": "Node.js",
    "express": "Express.js",
    "express.js": "Express.js",
    "mongo": "MongoDB",
    "mongodb": "MongoDB",
    "js": "JavaScript",
    "ts": "TypeScript",
    "react hooks": "React",
    "react components": "React",
    "crud apis": "FastAPI",
    "routing": "FastAPI",
}


def learning_skill(skill: str, role: str, importance: str = "") -> dict[str, Any]:
    name = re.sub(r"\s+", " ", str(skill)).strip()[:80]
    key = name.lower()
    advanced_terms = {"kubernetes", "deep learning", "mlops", "vector databases", "tensorflow", "model evaluation"}
    beginner_terms = {"html", "css", "git", "sql", "python", "excel", "communication", "projects"}
    inferred = "Advanced" if key in advanced_terms else "Beginner" if key in beginner_terms else "Intermediate"
    difficulty, duration = SKILL_LEARNING_META.get(
        key,
        (inferred, "24 Hours" if inferred == "Advanced" else "12 Hours" if inferred == "Intermediate" else "10 Hours"),
    )
    query = quote_plus(name)
    official = SKILL_OFFICIAL_RESOURCES.get(
        key,
        f"https://learn.microsoft.com/en-us/search/?terms={query}",
    )
    prerequisites = (
        ["Comfort with the skill's core concepts", "One small related project"]
        if difficulty == "Advanced"
        else ["Basic programming fundamentals"]
        if difficulty == "Intermediate"
        else ["No specialized prerequisite; start with the fundamentals"]
    )
    return {
        "name": name,
        "skill": name,
        "importance": importance or f"{name} is an important capability for practical {role} work and interviews.",
        "why": importance or f"{name} is an important capability for practical {role} work and interviews.",
        "difficulty": difficulty,
        "duration": duration,
        "prerequisites": prerequisites,
        "resources": {
            "official": official,
            "youtube": f"https://www.youtube.com/results?search_query=freeCodeCamp+{query}+full+course",
            "course": f"https://www.freecodecamp.org/news/search/?query={query}",
            "practice": f"https://github.com/search?q={query}+practice&type=repositories",
        },
    }


def normalize_skill_name(skill: str) -> str:
    name = re.sub(r"\s+", " ", str(skill)).strip()
    key = re.sub(r"[^a-z0-9+#. ]", "", name.lower()).strip()
    return SKILL_ALIASES.get(key, name)


def roadmap_resource(skill: str, resource_type: str) -> dict[str, str]:
    name = normalize_skill_name(skill)
    key = name.lower()
    query = quote_plus(name)
    overrides = ROADMAP_RESOURCE_OVERRIDES.get(key, {})
    official = overrides.get("official") or SKILL_OFFICIAL_RESOURCES.get(key) or f"https://learn.microsoft.com/en-us/search/?terms={query}"
    urls = {
        "official": official,
        "youtube": overrides.get("youtube") or f"https://www.youtube.com/results?search_query={query}+Tutorial",
        "course": overrides.get("course") or f"https://www.freecodecamp.org/news/search/?query={query}",
        "practice": overrides.get("practice") or f"https://github.com/search?q={query}+practice&type=repositories",
        "cheat_sheet": overrides.get("cheat_sheet") or f"https://learn.microsoft.com/en-us/search/?terms={query}+cheat+sheet",
    }
    labels = {
        "official": "Official Documentation",
        "youtube": "Complete Course",
        "course": "Free Course",
        "practice": "Practice",
        "cheat_sheet": "Cheat Sheet",
    }
    descriptions = {
        "official": f"Learn core {name} concepts, setup, APIs, and best practices from the source.",
        "youtube": f"Follow a guided {name} tutorial and build along from fundamentals to project work.",
        "course": f"Use a structured free course to turn {name} concepts into portfolio-ready skills.",
        "practice": f"Apply {name} through hands-on exercises, templates, and small build challenges.",
        "cheat_sheet": f"Review essential {name} syntax, commands, patterns, and quick reminders.",
    }
    providers = {
        "official": "Official Docs",
        "youtube": "YouTube",
        "course": "FreeCodeCamp / Free Course",
        "practice": "Practice Lab",
        "cheat_sheet": "Quick Revision",
    }
    estimated = {
        "official": "2 Hours",
        "youtube": "6-8 Hours",
        "course": "4-6 Hours",
        "practice": "3 Hours",
        "cheat_sheet": "30 Minutes",
    }
    actions = {
        "official": "Open Resource",
        "youtube": "Watch Video",
        "course": "Start Course",
        "practice": "Start Practice",
        "cheat_sheet": "Open Cheat Sheet",
    }
    return {
        "title": f"{name} {labels.get(resource_type, resource_type.replace('_', ' ').title())}",
        "description": descriptions.get(resource_type, f"Learn {name} with a trusted external resource."),
        "provider": providers.get(resource_type, "Learning Resource"),
        "badge": labels.get(resource_type, resource_type.replace("_", " ").title()),
        "estimated_time": estimated.get(resource_type, "1 Hour"),
        "action_label": actions.get(resource_type, "Open Resource"),
        "type": resource_type,
        "url": urls[resource_type],
        "skill": name,
    }


def infer_week_skills(week: dict[str, Any], missing_skills: list[str], required_skills: list[str]) -> list[str]:
    skill_items = week.get("skills_to_learn") or []
    if not isinstance(skill_items, list):
        skill_items = [str(skill_items)]
    task_items = week.get("tasks") or []
    if not isinstance(task_items, list):
        task_items = [str(task_items)]
    parts = [
        str(week.get("goal") or ""),
        str(week.get("phase") or ""),
        " ".join(str(item) for item in skill_items),
        " ".join(str(item) for item in task_items),
        str(week.get("mini_project") or ""),
        str(week.get("practice_goal") or ""),
    ]
    text = " ".join(parts).lower()
    candidates = clean_list([*missing_skills, *required_skills, *ROLE_SKILLS.get("frontend", []), *ROLE_SKILLS.get("backend", []), *ROLE_SKILLS.get("ai", [])], 80)
    found: list[str] = []
    for candidate in candidates:
        normalized = normalize_skill_name(candidate)
        aliases = {candidate.lower(), normalized.lower()}
        aliases.update(alias for alias, canonical in SKILL_ALIASES.items() if canonical.lower() == normalized.lower())
        if any(re.search(rf"(?<![a-z0-9]){re.escape(alias)}(?![a-z0-9])", text) for alias in aliases):
            found.append(normalized)
    if not found:
        found = [normalize_skill_name(skill) for skill in (skill_items or missing_skills[:2] or required_skills[:2])]
    return clean_list(found, 80)[:3]


def enrich_roadmap_resources(weeks: list[dict[str, Any]], role: str, missing_skills: list[str], required_skills: list[str]) -> list[dict[str, Any]]:
    enriched_weeks: list[dict[str, Any]] = []
    for index, week in enumerate(weeks, start=1):
        week_data = dict(week)
        week_data["week"] = int(week_data.get("week") or index)
        week_data["difficulty"] = week_data.get("difficulty") or ("Intermediate" if index <= 3 else "Advanced")
        skills = infer_week_skills(week_data, missing_skills, required_skills)
        existing_skills = week_data.get("skills_to_learn") or []
        if not isinstance(existing_skills, list):
            existing_skills = [str(existing_skills)]
        week_data["skills_to_learn"] = clean_list([*existing_skills, *skills], 80)
        resources: list[dict[str, str]] = []
        for skill in skills:
            for resource_type in ["official", "youtube", "course", "practice", "cheat_sheet"]:
                resources.append(roadmap_resource(skill, resource_type))
        seen: set[tuple[str, str]] = set()
        merged: list[dict[str, str]] = []
        for resource in [*week_data.get("resources", []), *resources]:
            if not isinstance(resource, dict) or not resource.get("url"):
                continue
            title = str(resource.get("title") or "").strip()
            resource_type = str(resource.get("type") or "resource").strip().lower()
            key = (resource_type, str(resource.get("url")))
            if key in seen:
                continue
            seen.add(key)
            skill = str(resource.get("skill") or "").strip()
            hydrated = roadmap_resource(skill, resource_type) if skill else {}
            merged.append(
                {
                    **hydrated,
                    "title": title or hydrated.get("title") or resource_type.title(),
                    "description": str(resource.get("description") or hydrated.get("description") or "Trusted learning resource for this week's roadmap."),
                    "provider": str(resource.get("provider") or hydrated.get("provider") or "Learning Resource"),
                    "badge": str(resource.get("badge") or hydrated.get("badge") or resource_type.replace("_", " ").title()),
                    "estimated_time": str(resource.get("estimated_time") or hydrated.get("estimated_time") or "1 Hour"),
                    "action_label": str(resource.get("action_label") or hydrated.get("action_label") or "Open Resource"),
                    "type": resource_type,
                    "url": str(resource["url"]),
                    "skill": skill or str(hydrated.get("skill") or ""),
                }
            )
        week_data["resources"] = merged
        enriched_weeks.append(week_data)
    return enriched_weeks


def build_personalized_roadmap(context: dict[str, Any], role: str, missing_skills: list[str], readiness: int, detected_skills: list[str]) -> dict[str, Any]:
    required = role_required_skills(role)
    focus = missing_skills or [skill for skill in required if skill.lower() not in {item.lower() for item in detected_skills}]
    if not focus:
        focus = required[:4]
    weeks: list[dict[str, Any]] = []
    for index in range(4):
        skills = focus[index * 2 : index * 2 + 2] or focus[:2]
        goal = f"Strengthen {' & '.join(skills)}" if skills else f"Build {role} readiness"
        weeks.append(
            {
                "week": index + 1,
                "goal": goal,
                "skills_to_learn": skills,
                "tasks": [f"Learn {skill} fundamentals" for skill in skills] + ([f"Build one {role} practice feature"] if index == 0 else ["Document progress in GitHub"]),
                "estimated_hours": 8 + index * 2,
                "status": "Not started",
                "progress": 0,
                "mini_project": f"Add {' and '.join(skills) if skills else 'core'} evidence to a portfolio project.",
                "practice_goal": f"Complete focused practice for {' and '.join(skills) if skills else role}.",
            }
        )
    weeks = enrich_roadmap_resources(weeks, role, focus, required)
    return {
        "current_readiness": readiness,
        "expected_readiness_after_completion": min(95, max(readiness + 20, readiness)),
        "experience_level": "Beginner" if readiness < 45 else "Intermediate" if readiness < 75 else "Advanced",
        "recommended_before_applying": [f"Add evidence of {skill} to a documented project." for skill in focus[:4]],
        "milestones": [{"title": f"Learn {skill}", "completed": skill.lower() in {item.lower() for item in detected_skills}} for skill in focus[:6]],
        "weeks": weeks,
    }


PLATFORM_HOME: dict[str, str] = {
    "freeCodeCamp": "https://www.freecodecamp.org/learn",
    "Coursera": "https://www.coursera.org/",
    "edX": "https://www.edx.org/",
    "Udemy": "https://www.udemy.com/",
    "Udacity": "https://www.udacity.com/",
    "Google Cloud Skills Boost": "https://www.cloudskillsboost.google/",
    "Microsoft Learn": "https://learn.microsoft.com/en-us/training/",
    "AWS Skill Builder": "https://skillbuilder.aws/",
    "MongoDB University": "https://learn.mongodb.com/",
    "GitHub Skills": "https://skills.github.com/",
    "Oracle University": "https://education.oracle.com/",
    "Cisco Skills for All": "https://skillsforall.com/",
    "Harvard CS50": "https://cs50.harvard.edu/",
    "Scrimba": "https://scrimba.com/",
    "React Official Docs": "https://react.dev/learn",
    "FastAPI Official Docs": "https://fastapi.tiangolo.com/",
}


COURSE_OVERRIDES: dict[str, dict[str, Any]] = {
    "html": {"title": "HTML & CSS Fundamentals", "platform": "freeCodeCamp", "duration": "12 Hours", "difficulty": "Beginner", "url": "https://www.freecodecamp.org/learn/2022/responsive-web-design/", "skills": ["HTML", "CSS", "Responsive Design"]},
    "css": {"title": "HTML & CSS Fundamentals", "platform": "freeCodeCamp", "duration": "12 Hours", "difficulty": "Beginner", "url": "https://www.freecodecamp.org/learn/2022/responsive-web-design/", "skills": ["HTML", "CSS", "Responsive Design"]},
    "javascript": {"title": "JavaScript Algorithms and Data Structures", "platform": "freeCodeCamp", "duration": "20 Hours", "difficulty": "Beginner", "url": "https://www.freecodecamp.org/learn/javascript-algorithms-and-data-structures-v8/", "skills": ["JavaScript", "Functions", "Data Structures"]},
    "typescript": {"title": "TypeScript Essentials", "platform": "Microsoft Learn", "duration": "10 Hours", "difficulty": "Intermediate", "url": "https://learn.microsoft.com/en-us/training/browse/?terms=TypeScript", "skills": ["Types", "Interfaces", "Type Safety"]},
    "react": {"title": "React Fundamentals", "platform": "Scrimba", "duration": "15 Hours", "difficulty": "Intermediate", "url": "https://scrimba.com/learn/learnreact", "skills": ["Components", "Hooks", "State Management"]},
    "fastapi": {"title": "FastAPI for Backend APIs", "platform": "FastAPI Official Docs", "duration": "8 Hours", "difficulty": "Intermediate", "url": "https://fastapi.tiangolo.com/tutorial/", "skills": ["Routing", "Validation", "CRUD APIs"]},
    "node.js": {"title": "Node.js & Express", "platform": "Coursera", "duration": "18 Hours", "difficulty": "Intermediate", "url": "https://www.coursera.org/search?query=node%20js%20express", "skills": ["Node.js", "Express", "REST APIs"]},
    "nodejs": {"title": "Node.js & Express", "platform": "Coursera", "duration": "18 Hours", "difficulty": "Intermediate", "url": "https://www.coursera.org/search?query=node%20js%20express", "skills": ["Node.js", "Express", "REST APIs"]},
    "express": {"title": "Node.js & Express", "platform": "Coursera", "duration": "18 Hours", "difficulty": "Intermediate", "url": "https://www.coursera.org/search?query=node%20js%20express", "skills": ["Node.js", "Express", "REST APIs"]},
    "mongodb": {"title": "MongoDB Developer Path", "platform": "MongoDB University", "duration": "12 Hours", "difficulty": "Intermediate", "url": "https://learn.mongodb.com/", "skills": ["MongoDB", "Schema Design", "CRUD"]},
    "sql": {"title": "SQL for Data and Backend Work", "platform": "Coursera", "duration": "14 Hours", "difficulty": "Beginner", "url": "https://www.coursera.org/search?query=sql", "skills": ["SQL", "Queries", "Joins"]},
    "git": {"title": "Git & GitHub", "platform": "GitHub Skills", "duration": "6 Hours", "difficulty": "Beginner", "url": "https://skills.github.com/", "skills": ["Git", "GitHub", "Pull Requests"]},
    "docker": {"title": "Docker Foundations", "platform": "Docker Official Docs", "duration": "10 Hours", "difficulty": "Intermediate", "url": "https://docs.docker.com/get-started/", "skills": ["Containers", "Images", "Deployment"]},
    "python": {"title": "Python Foundations", "platform": "freeCodeCamp", "duration": "20 Hours", "difficulty": "Beginner", "url": "https://www.freecodecamp.org/learn/scientific-computing-with-python/", "skills": ["Python", "Functions", "Automation"]},
    "machine learning": {"title": "Machine Learning Foundations", "platform": "Coursera", "duration": "24 Hours", "difficulty": "Intermediate", "url": "https://www.coursera.org/search?query=machine%20learning", "skills": ["ML", "Model Training", "Evaluation"]},
    "cloud deployment": {"title": "Cloud Deployment Fundamentals", "platform": "AWS Skill Builder", "duration": "12 Hours", "difficulty": "Intermediate", "url": "https://skillbuilder.aws/", "skills": ["Cloud", "Deployment", "Architecture"]},
}


def course_id(title: str, platform: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", f"{platform}-{title}".lower()).strip("-")


def platform_logo(platform: str, url: str) -> str:
    home = PLATFORM_HOME.get(platform, url)
    return f"https://www.google.com/s2/favicons?domain_url={quote_plus(home)}&sz=96"


def course_for_skill(skill: str, role: str, progress_state: dict[str, Any]) -> dict[str, Any]:
    name = normalize_skill_name(skill)
    key = name.lower()
    template = COURSE_OVERRIDES.get(key) or {
        "title": f"{name} Career Track",
        "platform": "Microsoft Learn",
        "duration": "10 Hours",
        "difficulty": "Intermediate",
        "url": f"https://learn.microsoft.com/en-us/search/?terms={quote_plus(name)}",
        "skills": [name, f"{role} Application", "Portfolio Practice"],
    }
    platform = str(template["platform"])
    title = str(template["title"])
    cid = course_id(title, platform)
    saved = progress_state.get(cid, {}) if isinstance(progress_state, dict) else {}
    progress = max(0, min(100, int(saved.get("progress", 0))))
    status = "Completed" if progress >= 100 else "In Progress" if progress > 0 else "Not Started"
    certificate = platform not in {"React Official Docs", "FastAPI Official Docs", "Docker Official Docs"}
    certificate_earned = bool(saved.get("certificate_earned") or (certificate and progress >= 100))
    return {
        "id": cid,
        "title": title,
        "platform": platform,
        "platformLogo": platform_logo(platform, str(template["url"])),
        "difficulty": str(template["difficulty"]),
        "duration": str(template["duration"]),
        "rating": float(template.get("rating", 4.7)),
        "certificate": bool(template.get("certificate", certificate)),
        "certificateAvailable": bool(template.get("certificate", certificate)),
        "certificateEarned": certificate_earned,
        "estimatedCompletionTime": str(template.get("estimatedCompletionTime", template["duration"])),
        "skills": clean_list(template.get("skills", [name]), 80),
        "progress": progress,
        "status": status,
        "url": str(template["url"]),
        "paid": bool(template.get("paid", platform in {"Coursera", "Udemy", "Udacity", "edX"})),
        "recommendationReason": f"{name} is a priority skill for {role} based on your profile, roadmap, or skill gap.",
    }


def catalog_role(role: str) -> str:
    requested = re.sub(r"\s+", " ", role).strip()
    aliases = {"ai engineer": "AI/ML Engineer", "ml engineer": "AI/ML Engineer", "software developer": "Software Engineer"}
    requested = aliases.get(requested.lower(), requested)
    return next((name for name in ROLE_COURSES if name.lower() == requested.lower()), "Software Engineer")


def recommended_courses(user: dict[str, Any], requested_role: str | None = None) -> dict[str, Any]:
    selected_role = requested_role or user.get("target_role") or "Software Engineer"
    role = catalog_role(str(selected_role))
    latest_gap = user.get("latest_skill_gap", {}) or {}
    missing_raw = latest_gap.get("missing_skills") or []
    missing: list[str] = []
    for item in missing_raw:
        name = item.get("name") or item.get("skill") if isinstance(item, dict) else item
        value = str(name).strip()
        if value:
            missing.append(value)
    missing_set = {skill.lower() for skill in clean_list(missing, 80)}
    current = {skill.lower() for skill in clean_list(user.get("skills", []), 80)}
    progress_state = user.get("course_progress", {})
    courses: list[dict[str, Any]] = []
    for stored in store.courses_for_role(role):
        course = dict(stored)
        saved = progress_state.get(course["id"], {}) if isinstance(progress_state, dict) else {}
        progress = max(0, min(100, int(saved.get("progress", 0))))
        skill_keys = {str(skill).lower() for skill in course.get("skills", [])}
        title_words = set(str(course.get("title", "")).lower().split())
        is_missing = bool(missing_set & (skill_keys | title_words))
        is_new_skill = not bool(current & (skill_keys | title_words))
        course.update({
            "platformLogo": platform_logo(str(course["platform"]), str(course["url"])),
            "certificateAvailable": bool(course.get("certificate")),
            "certificateEarned": bool(saved.get("certificate_earned") and course.get("certificate")),
            "progress": progress,
            "status": "Completed" if progress >= 100 else "In Progress" if progress else "Not Started",
            "recommendedForYou": is_missing or is_new_skill,
            "recommendationReason": (f"Your latest resume analysis identifies {course['title']} as a missing skill." if is_missing else f"Builds a priority skill for {role} based on your profile and ATS score ({int(user.get('resume_score') or 0)}%)."),
        })
        courses.append(course)
    courses.sort(key=lambda item: (not item["recommendedForYou"], item["roadmapOrder"]))
    completed = [course for course in courses if course["status"] == "Completed"]
    next_course = next((course for course in courses if course["status"] != "Completed"), courses[0] if courses else None)
    return {
        "role": str(selected_role),
        "targetRole": str(selected_role),
        "courses": courses,
        "completed": len(completed),
        "certificatesEarned": sum(1 for course in courses if course.get("certificateEarned")),
        "recommendedNextCourse": None if not next_course else {
            "title": next_course["title"],
            "platform": next_course["platform"],
            "reason": next_course["recommendationReason"],
        },
    }


def enrich_career_gps(result: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    current = {str(skill).strip().lower() for skill in context.get("skills", [])}
    for role_data in result.get("target_roles") or []:
        role = str(role_data.get("role") or context.get("target_role") or "the selected role")
        raw_missing = role_data.get("missing_skills") or []
        required = role_data.get("required_skills") or role_required_skills(role)
        known: dict[str, dict[str, Any]] = {}
        for item in raw_missing:
            source = item if isinstance(item, dict) else {"skill": str(item)}
            name = str(source.get("name") or source.get("skill") or "").strip()
            if not name:
                continue
            enriched = learning_skill(name, role, str(source.get("importance") or source.get("why") or ""))
            enriched["priority"] = source.get("priority") or "Medium"
            known[name.lower()] = enriched
        for skill in required:
            name = str(skill).strip()
            if name and name.lower() not in current and name.lower() not in known:
                enriched = learning_skill(name, role)
                enriched["priority"] = "Medium"
                known[name.lower()] = enriched
        role_data["missing_skills"] = list(known.values())
        missing_names = [item["name"] for item in role_data["missing_skills"]]
        role_data["learning_roadmap"] = enrich_roadmap_resources(role_data.get("learning_roadmap") or [], role, missing_names, required)
    return result


def role_required_skills(role: str) -> list[str]:
    lowered = role.lower()
    for key, skills in ROLE_SKILLS.items():
        if key in lowered:
            return skills
    return ["Python", "Problem Solving", "Git", "APIs", "SQL", "Communication", "Projects", "Testing", "Deployment"]


def skill_readiness(current_skills: list[str], required_skills: list[str]) -> tuple[int, list[str], list[str]]:
    current = clean_list(current_skills, 80)
    current_keys = {item.lower() for item in current}
    matched = [skill for skill in required_skills if skill.lower() in current_keys]
    missing = [skill for skill in required_skills if skill.lower() not in current_keys]
    score = round((len(matched) / max(1, len(required_skills))) * 100)
    return score, matched, missing


def local_career_gps(context: dict[str, Any], target_roles: list[str]) -> dict[str, Any]:
    roles = target_roles or [context.get("target_role") or "Software Developer"]
    results: list[dict[str, Any]] = []
    for role in roles:
        required = role_required_skills(role)
        readiness, matched, missing = skill_readiness(context.get("skills", []), required)
        priority = [
            {"skill": skill, "priority": "High" if index < 3 else "Medium", "why": f"{skill} is commonly expected for {role} roles."}
            for index, skill in enumerate(missing)
        ]
        results.append(
            {
                "role": role,
                "job_readiness_percentage": readiness,
                "current_skills": matched or context.get("skills", []),
                "required_skills": required,
                "missing_skills": priority,
                "learning_roadmap": [
                    {"phase": "Foundation", "duration": "1-2 weeks", "tasks": [f"Strengthen {skill}" for skill in missing[:2]] or ["Revise fundamentals", "Build one small practice task"]},
                    {"phase": "Projects", "duration": "2-4 weeks", "tasks": [f"Build a {role} portfolio project", "Add README, screenshots, and deployment link"]},
                    {"phase": "Interview prep", "duration": "1-2 weeks", "tasks": ["Practice role-specific questions", "Prepare project explanations", "Review common mistakes"]},
                ],
                "recommended_projects": [
                    f"{role} portfolio project with authentication and API integration",
                    "Dashboard or tracker that solves a real user problem",
                    "Polished GitHub project with tests and deployment notes",
                ],
                "certifications": ["FreeCodeCamp", "Coursera or Google career certificate", "Role-specific cloud fundamentals"],
                "free_resources": [
                    {"title": "MDN Web Docs", "provider": "Mozilla", "url": "https://developer.mozilla.org/"},
                    {"title": "freeCodeCamp", "provider": "freeCodeCamp", "url": "https://www.freecodecamp.org/"},
                ],
                "estimated_time_to_job_ready": "6-10 weeks with consistent project work",
            }
        )
    overall = round(sum(item["job_readiness_percentage"] for item in results) / len(results))
    return {"overall_readiness": overall, "target_roles": results, "cached": False}


def local_skill_gap(context: dict[str, Any]) -> dict[str, Any]:
    role = context.get("target_role") or "Software Developer"
    required = role_required_skills(role)
    readiness, matched, missing = skill_readiness(context.get("skills", []), required)
    priority = [{"skill": skill, "priority": "High" if index < 3 else "Medium", "reason": f"Needed for practical {role} work and interviews."} for index, skill in enumerate(missing)]
    return {
        "target_career": role,
        "job_readiness_percentage": readiness,
        "current_skills": matched or context.get("skills", []),
        "required_skills": required,
        "missing_skills": missing,
        "beginner_skills": missing[:3],
        "intermediate_skills": missing[3:7],
        "advanced_skills": missing[7:],
        "priority_order": priority,
        "recommended_learning_sequence": missing[:6] or ["Build a stronger portfolio project", "Practice interviews", "Deploy your best project"],
        "project_suggestions": [f"Build and deploy a {role} project", "Create a GitHub README case study", "Add tests and error handling"],
        "interview_topics": required[:6],
        "coding_practice_suggestions": ["Arrays and strings", "API design", "Debugging", "Project explanation practice"],
        "estimated_completion_timeline": "6-10 weeks",
        "cached": False,
    }


def local_jobs(context: dict[str, Any]) -> dict[str, Any]:
    role = context.get("target_role") or "Software Developer"
    required = role_required_skills(role)
    readiness, _matched, missing = skill_readiness(context.get("skills", []), required)
    titles = [role, f"Junior {role}", f"{role} Intern"]
    jobs = [
        {
            "title": title,
            "match_percentage": max(35, min(92, readiness - index * 8 + 12)),
            "missing_skills": missing[:5],
            "required_skills": required,
            "salary_range": "Market dependent",
            "recommended_improvements": ["Finish one deployed project", "Improve GitHub README quality", "Practice role-specific interview questions"],
            "why_it_matches": f"This is a realistic search target based on your current skills and target role: {role}.",
            "job_search_keywords": [role, "junior", "remote", "internship"],
        }
        for index, title in enumerate(titles)
    ]
    return {"target_role": role, "jobs": jobs, "cached": False}


def local_mentor_reply(message: str, user: dict[str, Any]) -> str:
    context = profile_context(user)
    role = context.get("target_role") or "your target role"
    required = role_required_skills(role)
    readiness, matched, missing = skill_readiness(context.get("skills", []), required)
    next_steps = missing[:3] or ["polish your best project", "practice interviews", "apply with targeted keywords"]
    lowered = message.strip().lower()
    if lowered in {"hi", "hii", "hello", "hey", "hy"}:
        return (
            f"Hey {user.get('name', 'there')}! I can help you plan your path for {role}. "
            f"Right now your saved profile shows about {readiness}% readiness. "
            f"Your next best focus areas are {', '.join(next_steps)}. Ask me about a roadmap, resume, projects, interview prep, or job search."
        )
    if "profile" in lowered or "about me" in lowered:
        gap = user.get("latest_skill_gap", {}) or {}
        github_score = int((user.get("github_analysis", {}) or {}).get("score") or 0)
        return (
            f"{user.get('name', 'You')}, you're working toward {role} with {readiness}% skill readiness. "
            f"Your current skills are {', '.join(context.get('skills', [])) or 'not filled in yet'}, and your strongest matches are "
            f"{', '.join(matched) or 'still being assessed'}. Your resume score is {int(user.get('resume_score') or 0)}% "
            f"and your GitHub score is {github_score}%. I'd make {next_steps[0]} your next goal. Keep going—you've already made good progress."
        )
    if "job ready" in lowered or "job-ready" in lowered or "ready for" in lowered:
        return (
            f"You're at roughly {readiness}% readiness for {role}. You already cover {', '.join(matched) or 'a few foundations'}, "
            f"while {', '.join(missing[:4]) or 'portfolio polish and interview practice'} still needs attention. "
            f"Give those gaps a focused 6–10 weeks, with one role-specific project and weekly interview practice. "
            "I believe you're ready for the next step."
        )
    if "resume" in lowered:
        score = int(user.get("resume_score") or 0)
        return (
            f"Your resume is currently scoring {score}%. For a {role} application, weave in evidence for "
            f"{', '.join(missing[:4]) or 'the skills already in your profile'}, strengthen project outcomes with numbers, "
            "and make each bullet show what you built, how you built it, and the result. Keep going—you’re getting closer."
        )
    if "github" in lowered:
        github = user.get("github_analysis", {}) or {}
        return (
            f"Your GitHub score is {int(github.get('score') or 0)}%. For {role}, pin the repositories that best prove "
            f"{', '.join(matched[:3]) or 'your core technical skills'}, add setup steps and screenshots to every README, "
            "and keep commits consistent while you finish one substantial project. You're getting closer every week 🚀"
        )
    if "python" in lowered:
        return (
            "Start Python with the basics first: variables, loops, functions, lists, dictionaries, file handling, and error handling. "
            "Then build 2 small projects: a CLI task tracker and a simple API using FastAPI. "
            f"Since your target role is {role}, connect Python practice to real portfolio work and push every project to GitHub with a clear README."
        )
    return (
        f"You're working toward {role}, with roughly {readiness}% readiness. Your current strengths include "
        f"{', '.join(matched) if matched else 'the foundations in your saved profile'}. I'd focus next on {', '.join(next_steps)}. "
        "Build one small proof project, "
        "write a clear README, and practice explaining what problem it solves, what stack you used, and what you would improve next."
    )


def local_resume_analysis(text: str, role: str, rough: dict[str, list[str]]) -> dict[str, Any]:
    required = role_required_skills(role)
    found_skills = [skill for skill in required if re.search(rf"\b{re.escape(skill)}\b", text, flags=re.I)]
    missing = [skill for skill in required if skill not in found_skills]
    has_projects = bool(rough.get("projects")) or "project" in text.lower()
    has_experience = bool(rough.get("experience")) or "experience" in text.lower()
    has_education = bool(rough.get("education")) or "education" in text.lower()
    score = 35 + len(found_skills) * 4 + (10 if has_projects else 0) + (10 if has_experience else 0) + (5 if has_education else 0)
    score = max(20, min(92, score))
    return {
        "ats_score": score,
        "score_breakdown": {
            "formatting": 70 if len(text.splitlines()) > 8 else 45,
            "keywords": max(20, min(90, len(found_skills) * 10)),
            "experience": 75 if has_experience else 35,
            "skills": max(25, min(90, len(found_skills) * 10)),
            "projects": 80 if has_projects else 35,
            "education": 75 if has_education else 45,
        },
        "score_explanation": f"Resume checked for {role}. Improve score by adding measurable project bullets, required keywords, and clear impact statements.",
        "detected": {**rough, "skills": found_skills},
        "strengths": ["Readable resume text extracted", "Some role signals are present"] if found_skills else ["Readable resume text extracted"],
        "weaknesses": ["Add more measurable achievements", "Use role-specific keywords from the job description"],
        "missing_skills": missing[:8],
        "missing_keywords": missing[:8],
        "suggestions": ["Add numbers to project and experience bullets", "Group skills by category", "Add deployment/GitHub links for projects"],
        "improved_resume_points": [
            f"Built a {role} project using relevant tools and documented setup, features, and outcomes.",
            "Improved project quality by adding error handling, responsive UI, and clear README documentation.",
        ],
        "ats_friendly_bullets": [
            "Developed and deployed a production-style project with authentication, API integration, and clean documentation.",
            "Applied Git workflow and testing practices to improve code reliability and maintainability.",
        ],
        "cached": False,
    }


def local_interview_question(context: dict[str, Any], round_name: str, language: str, question_id: str) -> dict[str, Any]:
    role = context.get("target_role") or "Software Developer"
    if round_name.lower() == "coding":
        text = f"In {language}, write a function to solve a common array/string problem and explain its time complexity."
        focus = f"{language} problem solving"
    elif "technical" in round_name.lower():
        text = f"Explain one {role} project you would build, including architecture, APIs, data flow, and testing."
        focus = "technical depth"
    else:
        text = f"Tell me about yourself and why you are preparing for a {role} role."
        focus = "communication"
    return {
        "question": {
            "id": question_id,
            "text": text,
            "ideal_signals": ["clear structure", "specific examples", "honest tradeoffs", "role-relevant details"],
            "difficulty": "medium",
            "focus_area": focus,
        }
    }


def local_interview_feedback(payload: InterviewAnswerRequest) -> dict[str, Any]:
    words = payload.answer.split()
    length_score = min(90, max(35, len(words) * 3))
    has_example = any(item in payload.answer.lower() for item in ["project", "built", "used", "implemented", "because"])
    relevance = 80 if has_example else 55
    overall = round((length_score * 0.45) + (relevance * 0.35) + 60 * 0.2)
    return {
        "overall": overall,
        "criteria": {"technical": relevance, "communication": length_score, "confidence": 60, "relevance": relevance, "completeness": length_score},
        "strengths": ["You attempted the answer clearly"] + (["You included practical context"] if has_example else []),
        "weaknesses": ["Add more specific examples", "Structure the answer with situation, action, result"],
        "mistakes": [] if has_example else ["Answer is too generic"],
        "suggestions": ["Use one real project example", "Mention tools used", "Explain the result or learning"],
        "ideal_answer": "A strong answer gives context, explains your exact contribution, names the tools used, and ends with measurable impact or learning.",
        "practice_plan": ["Prepare 3 project stories", "Practice 2-minute answers", "Review common role-specific questions"],
    }


def local_github_analysis(username: str, repos: list[dict[str, Any]], profile: dict[str, Any]) -> dict[str, Any]:
    languages = sorted({repo.get("language") for repo in repos if repo.get("language")})
    role = profile.get("target_role") or "Software Developer"
    score = min(90, 35 + len(repos[:10]) * 4 + len(languages) * 5)
    repo_analysis = [
        {"name": repo.get("name"), "feedback": "Add a strong README, screenshots, setup steps, and clear feature list."}
        for repo in repos[:6]
    ]
    return {
        "score": score,
        "languages_used": languages,
        "tech_stack": languages,
        "contribution_quality": "Improve consistency with frequent commits and clearer project history.",
        "project_quality": "Focus on complete, deployed projects with documentation and tests.",
        "repository_analysis": repo_analysis,
        "missing_portfolio_projects": [f"One polished {role} project", "A project with API integration", "A project with tests and deployment"],
        "recommended_projects": [f"{role} portfolio app", "Dashboard or tracker app", "API-backed full-stack project"],
        "coding_consistency": "Keep commits small and regular with meaningful messages.",
        "portfolio_feedback": f"Your GitHub has {len(repos)} visible repositories. Make the best 2-3 repos recruiter-ready.",
        "recommendations": ["Pin best repositories", "Write better READMEs", "Add screenshots", "Deploy projects"],
        "repos": repos[:8],
        "cached": False,
    }


def github_username(identifier: str) -> str:
    value = identifier.strip().rstrip("/")
    match = re.fullmatch(
        r"(?:https?://)?(?:www\.)?github\.com/([^/?#]+)(?:[/?#].*)?",
        value,
        flags=re.IGNORECASE,
    )
    username = (match.group(1) if match else value).lstrip("@").strip()
    if not re.fullmatch(r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,37}[A-Za-z0-9])?", username):
        raise HTTPException(status_code=400, detail="Enter a valid GitHub username or profile URL.")
    return username


def github_headers() -> dict[str, str]:
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "ai-career-copilot",
        "X-GitHub-Api-Version": "2022-11-28",
    }
    if GITHUB_TOKEN:
        headers["Authorization"] = f"Bearer {GITHUB_TOKEN}"
    return headers


GITHUB_SESSION = requests.Session()
GITHUB_SESSION.mount(
    "https://api.github.com",
    HTTPAdapter(
        max_retries=Retry(
            total=2,
            connect=2,
            read=2,
            status=2,
            backoff_factor=0.5,
            status_forcelist=(429, 500, 502, 503, 504),
            allowed_methods=frozenset({"GET"}),
            respect_retry_after_header=True,
        )
    ),
)


def github_get(url: str, **kwargs: Any) -> requests.Response:
    kwargs.setdefault("headers", github_headers())
    kwargs.setdefault("timeout", GITHUB_TIMEOUT)
    kwargs.setdefault("verify", REQUEST_VERIFY)
    return GITHUB_SESSION.get(url, **kwargs)


def readme_score(full_name: str) -> tuple[int, str]:
    try:
        response = github_get(
            f"https://api.github.com/repos/{full_name}/readme",
        )
        if response.status_code == 404:
            return 0, ""
        response.raise_for_status()
        encoded = response.json().get("content", "")
        text = base64.b64decode(encoded).decode("utf-8", errors="ignore") if encoded else ""
    except (requests.RequestException, ValueError):
        return 0, ""
    lowered = text.lower()
    score = min(35, len(text) // 80)
    score += 15 if re.search(r"^#{1,3}\s+", text, re.MULTILINE) else 0
    score += 15 if any(word in lowered for word in ["install", "getting started", "setup"]) else 0
    score += 15 if any(word in lowered for word in ["usage", "features", "demo"]) else 0
    score += 10 if any(word in lowered for word in ["license", "contributing"]) else 0
    score += 10 if any(word in lowered for word in ["![", "<img", "screenshot"]) else 0
    return min(100, score), text


REPOSITORY_DEVELOPMENT_SUMMARY = "This repository is currently under development. More project information will become available as additional documentation and source code are added."


def readme_summary(readme: str) -> str:
    if not readme.strip():
        return ""
    cleaned = re.sub(r"```.*?```", " ", readme, flags=re.S)
    cleaned = re.sub(r"!\[[^]]*\]\([^)]*\)|<[^>]+>|https?://\S+", " ", cleaned)
    paragraphs = []
    for block in re.split(r"\n\s*\n", cleaned):
        value = re.sub(r"[#>*_`|\[\]()]", " ", block)
        value = re.sub(r"\s+", " ", value).strip(" -:;")
        if len(value) >= 45 and not value.lower().startswith(("installation", "getting started", "usage", "license", "contents")):
            paragraphs.append(value)
        if sum(len(item) for item in paragraphs) >= 320:
            break
    summary = " ".join(paragraphs)[:520].rsplit(" ", 1)[0].strip()
    return summary if len(summary) >= 45 else ""


def repository_file_signals(repo: dict[str, Any]) -> dict[str, Any]:
    full_name = str(repo.get("full_name") or "")
    branch = str(repo.get("default_branch") or "main")
    if not full_name:
        return {"files": [], "manifests": [], "source_files": [], "file_excerpts": {}}
    try:
        response = github_get(f"https://api.github.com/repos/{full_name}/git/trees/{branch}", params={"recursive": "1"})
        if not response.ok:
            return {"files": [], "manifests": [], "source_files": [], "file_excerpts": {}}
        paths = [str(item.get("path") or "") for item in response.json().get("tree", []) if item.get("type") == "blob"][:500]
    except (requests.RequestException, ValueError):
        return {"files": [], "manifests": [], "source_files": [], "file_excerpts": {}}
    manifest_names = {"package.json", "requirements.txt", "pom.xml", "pyproject.toml", "build.gradle", "pubspec.yaml", "dockerfile"}
    manifests = [path for path in paths if Path(path).name.lower() in manifest_names]
    sources = [path for path in paths if Path(path).suffix.lower() in {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".dart", ".go", ".cpp", ".c"}]
    excerpts: dict[str, str] = {}
    for path in [*manifests[:2], *sources[:2]]:
        try:
            content_response = github_get(f"https://api.github.com/repos/{full_name}/contents/{quote(path, safe='/')}", params={"ref": branch})
            if not content_response.ok:
                continue
            encoded = content_response.json().get("content", "")
            content = base64.b64decode(encoded).decode("utf-8", errors="ignore") if encoded else ""
            if content.strip():
                excerpts[path] = content[:3500]
        except (requests.RequestException, ValueError):
            continue
    return {"files": paths[:80], "manifests": manifests[:12], "source_files": sources[:20], "file_excerpts": excerpts}


def professional_repository_summary(repo: dict[str, Any], readme: str, signals: dict[str, Any]) -> str:
    summary = readme_summary(readme)
    if summary:
        return summary
    name = str(repo.get("name") or "").replace("-", " ").replace("_", " ").strip()
    language = str(repo.get("language") or "").strip()
    topics = clean_list([str(item) for item in repo.get("topics", [])], 60)
    evidence = [language, *topics, *[Path(path).name for path in signals.get("manifests", [])]]
    evidence = [item for item in clean_list(evidence, 80) if item.lower() not in {"other", "none", "null"}]
    if evidence:
        stack = ", ".join(evidence[:5])
        return f"{name.title()} is a software project built with {stack}. Its repository structure indicates an actively developed implementation designed for practical use and continued feature expansion."
    return REPOSITORY_DEVELOPMENT_SUMMARY


def ai_repository_summaries(items: list[dict[str, Any]]) -> dict[str, str]:
    if not items:
        return {}
    schema = {"summaries": [{"name": "", "description": ""}]}
    try:
        result = groq_json(
            "You are a senior technical portfolio writer. Produce accurate, professional repository descriptions from evidence only. Never mention missing data, parsing, APIs, or uncertainty. Return JSON only.",
            f"Write a polished 2-3 sentence project summary for each repository. Explain its purpose, capabilities, and technology without inventing unsupported features. Evidence: {json.dumps(items)[:24000]}",
            schema,
            0.35,
        )
    except HTTPException:
        return {}
    return {str(item.get("name")): str(item.get("description") or "").strip() for item in result.get("summaries", []) if item.get("name") and len(str(item.get("description") or "").strip()) >= 45}


def github_portfolio_report(
    username: str,
    account: dict[str, Any],
    repos: list[dict[str, Any]],
    context: dict[str, Any],
) -> dict[str, Any]:
    analyzed: list[dict[str, Any]] = []
    language_counts: dict[str, int] = {}
    detected: set[str] = set()
    activity: dict[str, int] = {}
    summary_requests: list[dict[str, Any]] = []
    language_aliases = {
        "javascript": "JavaScript", "typescript": "TypeScript", "html": "HTML", "css": "CSS",
        "python": "Python", "java": "Java", "c++": "C++", "c#": "C#", "go": "Go",
        "ruby": "Ruby", "php": "PHP", "kotlin": "Kotlin", "swift": "Swift", "dart": "Dart",
    }
    if repos:
        detected.add("Git")

    for repo in repos[:12]:
        language = repo.get("language") or "Other"
        language_counts[language] = language_counts.get(language, 0) + 1
        if language.lower() in language_aliases:
            detected.add(language_aliases[language.lower()])
        topics = [str(item) for item in repo.get("topics", [])]
        searchable = " ".join([repo.get("name") or "", repo.get("description") or "", *topics]).lower()
        for needle, skill in {
            "react": "React", "fastapi": "FastAPI", "django": "Django", "flask": "Flask",
            "node": "Node.js", "sql": "SQL", "mongo": "MongoDB", "docker": "Docker",
            "api": "API Integration", "tailwind": "Tailwind CSS", "machine-learning": "Machine Learning",
            "tensorflow": "TensorFlow", "pytorch": "PyTorch", "pandas": "Pandas", "numpy": "NumPy",
        }.items():
            if needle in searchable:
                detected.add(skill)

        readme_quality, readme = readme_score(repo.get("full_name") or f"{username}/{repo.get('name', '')}")
        github_description = str(repo.get("description") or "").strip()
        readme_description = readme_summary(readme) if not github_description else ""
        signals = repository_file_signals(repo) if not github_description and not readme_description else {"files": [], "manifests": [], "source_files": []}
        description = github_description or readme_description or professional_repository_summary(repo, readme, signals)
        if not github_description and not readme_description and (signals.get("files") or repo.get("language") or topics):
            summary_requests.append({"name": repo.get("name"), "language": repo.get("language"), "topics": topics, **signals})
        description_quality = 100 if repo.get("description") else 20
        documentation_quality = round(readme_quality * 0.8 + description_quality * 0.2)
        complexity = min(100, 20 + min(35, int(repo.get("size") or 0) // 150) + len(topics) * 5)
        complexity += 10 if repo.get("language") else 0
        complexity += 10 if repo.get("homepage") or repo.get("has_pages") else 0
        complexity = min(100, complexity)
        organization = min(100, 35 + (20 if readme else 0) + (15 if repo.get("license") else 0) + min(30, int(repo.get("size") or 0) // 200))
        health = round((readme_quality + documentation_quality + complexity + organization) / 4)
        production = round(
            health * 0.65
            + (15 if repo.get("homepage") or repo.get("has_pages") else 0)
            + min(10, int(repo.get("stargazers_count") or 0) * 2)
            + (10 if repo.get("license") else 0)
        )
        pushed = repo.get("pushed_at") or repo.get("updated_at")
        if pushed:
            month = pushed[:7]
            activity[month] = activity.get(month, 0) + 1
        tech_stack = clean_list(([language] if language != "Other" else []) + topics, 40)
        analyzed.append({
            "name": repo.get("name"),
            "description": description,
            "description_source": "github" if github_description else "readme" if readme_description else "generated",
            "url": repo.get("html_url"),
            "stars": int(repo.get("stargazers_count") or 0),
            "forks": int(repo.get("forks_count") or 0),
            "topics": topics,
            "languages": [] if language == "Other" else [language],
            "tech_stack": tech_stack,
            "repository_health": health,
            "production_readiness": min(100, production),
            "readme_quality": readme_quality,
            "documentation_quality": documentation_quality,
            "project_complexity": complexity,
            "code_organization": organization,
            "deployment_status": "Deployed" if repo.get("homepage") or repo.get("has_pages") else "",
            "deployment_url": repo.get("homepage") or (f"https://{username}.github.io/{repo.get('name')}" if repo.get("has_pages") else ""),
            "last_updated": repo.get("pushed_at") or repo.get("updated_at"),
            "github_actions": any(path.startswith(".github/workflows/") for path in signals.get("files", [])),
            "license": (repo.get("license") or {}).get("spdx_id") or "",
        })

    generated_summaries = ai_repository_summaries(summary_requests)
    for repository in analyzed:
        generated = generated_summaries.get(str(repository.get("name")))
        if repository.get("description_source") == "generated" and generated:
            repository["description"] = generated

    repo_count = len(repos)
    stars = sum(int(repo.get("stargazers_count") or 0) for repo in repos)
    forks = sum(int(repo.get("forks_count") or 0) for repo in repos)
    average = lambda key: round(sum(item[key] for item in analyzed) / max(1, len(analyzed)))
    profile_fields = [account.get(key) for key in ["name", "bio", "company", "location", "blog"]]
    profile_points = sum(2 for value in profile_fields if value)
    average_health = average("repository_health")
    average_readme = average("readme_quality")
    average_complexity = average("project_complexity")
    average_docs = average("documentation_quality")
    recent_repos = sum(1 for repo in repos if (repo.get("pushed_at") or "")[:4] in {str(datetime.now().year), str(datetime.now().year - 1)})
    breakdown = {
        "Profile Completion": profile_points,
        "README Quality": round(average_readme * 0.15),
        "Repository Quality": round(average_health * 0.20),
        "Project Complexity": round(average_complexity * 0.20),
        "Commit Consistency": min(10, recent_repos * 2),
        "Code Diversity": min(10, len(language_counts) * 2),
        "Documentation": round(average_docs * 0.05),
        "Portfolio Website": 5 if account.get("blog") else 0,
        "Pinned Projects": min(5, sum(1 for item in analyzed if item["production_readiness"] >= 65) * 2),
        "Open Source Contribution": min(5, sum(1 for repo in repos if repo.get("fork")) * 2 + min(3, forks)),
    }
    score = min(100, round(sum(breakdown.values()) / 105 * 100))
    role = context.get("target_role") or "Software Developer"
    required = role_required_skills(role)
    readiness, matched, missing = skill_readiness(sorted(detected), required)
    best = sorted(analyzed, key=lambda item: (item["production_readiness"], item["stars"]), reverse=True)
    strengths = []
    if repo_count:
        strengths.append(f"{repo_count} public repositories demonstrate hands-on project work.")
    if len(language_counts) >= 2:
        strengths.append(f"Projects use {len(language_counts)} primary languages.")
    if average_readme >= 60:
        strengths.append("Repository documentation is generally recruiter-friendly.")
    weaknesses = []
    if average_readme < 60:
        weaknesses.append("Several repositories need stronger README files with setup, features, and screenshots.")
    if not account.get("blog"):
        weaknesses.append("No portfolio website is linked from the GitHub profile.")
    if not any(item["deployment_url"] for item in analyzed):
        weaknesses.append("No live project deployment was detected.")
    recommendations = []
    if missing:
        recommendations.append(f"Show {', '.join(missing[:3])} in a project tailored to {role} roles.")
    recommendations.extend(weaknesses[:2])
    if not recommendations:
        recommendations.append("Keep the strongest projects current and pin the most job-relevant repositories.")

    blog = str(account.get("blog") or "")
    if blog and not re.match(r"https?://", blog):
        blog = f"https://{blog}"
    return {
        "score": score,
        "rating": "Excellent portfolio" if score >= 80 else "Strong portfolio" if score >= 65 else "Developing portfolio" if score >= 45 else "Needs portfolio improvements",
        "profile": {
            "username": account.get("login") or username,
            "name": account.get("name") or account.get("login") or username,
            "avatar_url": account.get("avatar_url"),
            "url": account.get("html_url") or f"https://github.com/{username}",
            "bio": account.get("bio"),
            "followers": int(account.get("followers") or 0),
            "following": int(account.get("following") or 0),
            "public_repositories": int(account.get("public_repos") or repo_count),
            "portfolio_website": blog if "linkedin.com" not in blog.lower() else "",
            "linkedin": blog if "linkedin.com" in blog.lower() else "",
        },
        "totals": {"repositories": repo_count, "stars": stars, "forks": forks},
        "score_breakdown": breakdown,
        "career_readiness": {
            "percentage": readiness,
            "target_role": role,
            "detected_skills": matched,
            "missing_skills": missing,
            "recommended_before_applying": [f"Add evidence of {skill} to a documented project." for skill in missing[:4]],
        },
        "charts": {
            "language_usage": [{"name": name, "value": value} for name, value in sorted(language_counts.items(), key=lambda pair: pair[1], reverse=True)],
            "commit_frequency": [{"name": name, "value": value} for name, value in sorted(activity.items())],
            "skill_distribution": [{"name": name, "value": sum(name.lower() in " ".join(item["tech_stack"]).lower() for item in analyzed)} for name in sorted(detected)],
        },
        "detected_skills": sorted(detected),
        "strengths": strengths or ["The GitHub profile is active and available for portfolio development."],
        "weaknesses": weaknesses,
        "recommendations": recommendations,
        "best_projects": best[:4],
        "repository_analysis": analyzed,
        "repos": analyzed[:8],
        "cached": False,
        "partial": False,
    }


def extract_pdf_details(raw: bytes) -> tuple[str, list[dict[str, Any]], bool]:
    if PdfReader is None:
        raise HTTPException(status_code=500, detail="PDF parser is not installed.")
    try:
        reader = PdfReader(io.BytesIO(raw))
        page_text: list[str] = []
        layout: list[dict[str, Any]] = []
        for index, page in enumerate(reader.pages):
            if index == 0:
                height = float(page.mediabox.height or 0)

                def visit(text: str, _cm: Any, tm: Any, _font: Any, font_size: float) -> None:
                    value = re.sub(r"\s+", " ", text).strip()
                    top = height - float(tm[5] if len(tm) > 5 else 0)
                    if value and top <= height * 0.42:
                        layout.append({"text": value, "font_size": float(font_size or 0), "top": top})

                page_text.append(page.extract_text(visitor_text=visit) or "")
            else:
                page_text.append(page.extract_text() or "")
        text = "\n\f\n".join(page_text)
        used_ocr = False
        if len(re.sub(r"\s+", "", text)) < 80:
            ocr = ocr_pdf(raw)
            if ocr.strip():
                text = ocr
                used_ocr = True
        return text, layout, used_ocr
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not extract text from PDF: {exc}") from exc


def ocr_pdf(raw: bytes) -> str:
    """OCR image-only PDFs when PyMuPDF, Pillow, and Tesseract are available."""
    if fitz is None or pytesseract is None or Image is None:
        return ""
    try:
        document = fitz.open(stream=raw, filetype="pdf")
        pages: list[str] = []
        for page in document[:5]:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            pages.append(pytesseract.image_to_string(image))
        return "\n\f\n".join(pages)
    except Exception:
        return ""


def extract_pdf(raw: bytes) -> str:
    return extract_pdf_details(raw)[0]


def extract_docx(raw: bytes) -> str:
    if Document is None:
        raise HTTPException(status_code=500, detail="DOCX parser is not installed.")
    try:
        document = Document(io.BytesIO(raw))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not extract text from DOCX: {exc}") from exc


def extract_resume_text(filename: str, raw: bytes) -> str:
    return extract_resume_document(filename, raw)[0]


def extract_resume_document(filename: str, raw: bytes) -> tuple[str, list[dict[str, Any]], bool]:
    if len(raw) > 8 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Resume file must be 8MB or smaller.")
    suffix = Path(filename).suffix.lower()
    content_type_text = raw[:16].strip().startswith((b"{", b"Name", b"name"))
    layout: list[dict[str, Any]] = []
    used_ocr = False
    if suffix == ".pdf":
        text, layout, used_ocr = extract_pdf_details(raw)
    elif suffix == ".docx":
        text = extract_docx(raw)
    elif suffix in {".txt", ".md"} or content_type_text:
        text = raw.decode("utf-8", errors="ignore")
    else:
        raise HTTPException(status_code=400, detail="Upload a PDF, DOCX, TXT, or MD resume.")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise HTTPException(status_code=400, detail="Could not extract enough readable resume text.")
    return text[:MAX_RESUME_CHARS], layout, used_ocr


def rough_resume_sections(text: str) -> dict[str, list[str]]:
    headings = ["education", "experience", "projects", "skills", "certifications", "achievements"]
    sections = {heading: [] for heading in headings}
    current = ""
    for line in [item.strip(" -\t") for item in text.splitlines() if item.strip()]:
        key = re.sub(r"[^a-z]", "", line.lower())
        matched = next((heading for heading in headings if key == heading or key.startswith(heading)), "")
        if matched:
            current = matched
            continue
        if current and len(sections[current]) < 20:
            sections[current].append(line[:300])
    return sections


def llm_cached(user: dict[str, Any], name: str, payload: dict[str, Any], schema: dict[str, Any], prompt: str, temperature: float = 0.25) -> dict[str, Any]:
    key = cache_key(name, payload)
    hit = cached(user, key)
    if hit is not None:
        return {**hit, "cached": True}
    result = groq_json(
        "You are a production career intelligence engine. Be specific, current, personalized, and honest. Never invent uploaded facts. Return compact JSON only.",
        prompt,
        schema,
        temperature,
    )
    result["cached"] = False
    set_cached(user, key, result)
    store.upsert_user(user)
    return result


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {
        "ok": True,
        "database": "mongodb" if store.using_mongo else "local-json",
        "groq_configured": groq_configured(),
        "groq_model": GROQ_MODEL,
        "time": now_iso(),
    }


@app.post("/api/auth/register")
def register(payload: RegisterRequest) -> dict[str, Any]:
    if store.get_user(payload.email):
        raise HTTPException(status_code=409, detail="Email already registered")
    user = default_user(payload.name, payload.email, payload.password)
    store.upsert_user(user)
    return {"token": create_token(user["email"]), "user": public_user(user)}


@app.post("/api/auth/login")
def login(payload: LoginRequest) -> dict[str, Any]:
    user = store.get_user(payload.email)
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    return {"token": create_token(user["email"]), "user": public_user(user)}


@app.get("/api/me")
def me(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    return public_user(user)


@app.put("/api/profile")
def update_profile(payload: ProfileRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    user.update(payload.model_dump())
    store.upsert_user(user)
    return public_user(user)


@app.get("/api/dashboard")
def dashboard(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    resume_score = int(user.get("resume_score") or 0)
    github_score = int(user.get("github_analysis", {}).get("score") or 0)
    latest_gap = user.get("latest_skill_gap", {})
    readiness = int(latest_gap.get("job_readiness_percentage") or 0)
    success = round((resume_score * 0.35) + (github_score * 0.25) + (readiness * 0.4))
    missing = latest_gap.get("missing_skills", [])
    course_data = recommended_courses(user)
    return {
        "success": success,
        "skill_score": readiness,
        "resume_score": resume_score,
        "github_score": github_score,
        "xp": len(user.get("interview_sessions", [])) * 120 + len(user.get("skills", [])) * 45,
        "level": max(1, min(20, 1 + len(user.get("skills", [])) // 3)),
        "matched_skills": latest_gap.get("current_skills", user.get("skills", [])),
        "missing_skills": missing,
        "course_progress": {
            "completed": course_data["completed"],
            "total": len(course_data["courses"]),
            "certificates_earned": course_data["certificatesEarned"],
        },
    }


@app.get("/api/courses")
def courses(role: str = "", user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    return recommended_courses(user, role or None)


@app.get("/api/v1/course-recommendations")
def course_recommendations(role: str = "", user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    """Return the persisted catalog for a role, personalized for this user."""
    return recommended_courses(user, role or None)


@app.post("/api/courses/progress")
def update_course_progress(payload: CourseProgressRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    current = user.setdefault("course_progress", {})
    state = current.setdefault(payload.course_id, {})
    state["progress"] = payload.progress
    state["status"] = "Completed" if payload.progress >= 100 else "In Progress" if payload.progress > 0 else "Not Started"
    state["updated_at"] = now_iso()
    if payload.progress >= 100:
        state["completed_at"] = now_iso()
        state["certificate_earned"] = True
        certificates = user.setdefault("certificates", [])
        if not any(item.get("course_id") == payload.course_id for item in certificates):
            certificates.append({"course_id": payload.course_id, "earned_at": now_iso(), "source": "Course Recommendations"})
    store.upsert_user(user)
    return recommended_courses(user)


@app.post("/api/resume/analyze")
async def analyze_resume(
    file: UploadFile | None = File(default=None),
    target_role: str = Form(default=""),
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    if file:
        text, layout, used_ocr = extract_resume_document(file.filename or "resume", await file.read())
    else:
        text = user.get("resume_text", "")
        layout, used_ocr = [], False
        if not text:
            raise HTTPException(status_code=400, detail="Upload a resume before analyzing saved resume.")
    role = target_role.strip() or user.get("target_role", "")
    rough = rough_resume_sections(text)
    identity = extract_profile(text, layout)
    payload = {"resume_hash": hashlib.sha256(text.encode()).hexdigest(), "target_role": role}
    schema = {
        "ats_score": 0,
        "score_breakdown": {"formatting": 0, "keywords": 0, "experience": 0, "skills": 0, "projects": 0, "education": 0},
        "score_explanation": "",
        "detected": {"name": "", "education": [], "experience": [], "projects": [], "skills": [], "certifications": [], "achievements": []},
        "strengths": [],
        "weaknesses": [],
        "missing_skills": [],
        "missing_keywords": [],
        "suggestions": [],
        "improved_resume_points": [],
        "ats_friendly_bullets": [],
    }
    try:
        result = llm_cached(
            user,
            "resume_ats",
            payload,
            schema,
            f"Analyze this resume for ATS and role fit. Target role: {role}. Rough extracted sections: {json.dumps(rough)}. Resume text:\n{text}",
        )
    except HTTPException as exc:
        if not is_groq_auth_error(exc):
            raise
        result = local_resume_analysis(text, role, rough)
    score = int(result.get("ats_score") or result.get("score") or 0)
    user["resume_score"] = max(0, min(100, score))
    user["resume_text"] = text
    user["resume_sections"] = result.get("detected", rough)
    store.upsert_user(user)
    result["score"] = user["resume_score"]
    result["extracted_text_length"] = len(text)
    detected = result.get("detected") or {}
    detected.update(identity)
    detected.setdefault("education", rough.get("education", []))
    detected.setdefault("experience", rough.get("experience", []))
    detected.setdefault("skills", rough.get("skills", []))
    result["detected"] = detected
    result["profile"] = {
        "name": identity["name"],
        "confidence": identity["confidence"],
        "email": identity["email"],
        "phone": identity["phone"],
        "skills": detected.get("skills", []),
        "education": detected.get("education", []),
        "experience": detected.get("experience", []),
    }
    result["ocr_used"] = used_ocr
    user["resume_sections"] = detected
    store.upsert_user(user)
    detected_skills = clean_list(detected.get("skills", []) or rough.get("skills", []), 80)
    missing_skills = clean_list(result.get("missing_skills") or result.get("missing_keywords") or [], 80)
    result["personalized_roadmap"] = build_personalized_roadmap(
        profile_context(user),
        role,
        missing_skills,
        user["resume_score"],
        detected_skills,
    )
    return result


@app.post("/api/gps")
def career_gps(payload: SkillRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    user["skills"] = payload.skills
    if payload.career_interests:
        user["career_interests"] = payload.career_interests
    context = profile_context(user)
    target_roles = payload.target_roles or [user.get("target_role", "")]
    request_payload = {"context": context, "target_roles": target_roles}
    schema = {
        "overall_readiness": 0,
        "target_roles": [
            {
                "role": "",
                "job_readiness_percentage": 0,
                "current_skills": [],
                "required_skills": [],
                "missing_skills": [{"skill": "", "priority": "", "why": ""}],
                "learning_roadmap": [{"phase": "", "duration": "", "tasks": []}],
                "recommended_projects": [],
                "certifications": [],
                "free_resources": [{"title": "", "provider": "", "url": ""}],
                "estimated_time_to_job_ready": "",
            }
        ],
    }
    try:
        result = llm_cached(
            user,
            "career_gps",
            request_payload,
            schema,
            f"Create a personalized AI career roadmap from this user profile. Determine required skills from current industry expectations, not a fixed list. Profile JSON:\n{json.dumps(request_payload)}",
        )
    except HTTPException as exc:
        if not is_groq_auth_error(exc):
            raise
        result = local_career_gps(context, target_roles)
    result = enrich_career_gps(result, context)
    roles = result.get("target_roles") or []
    if roles:
        user["latest_skill_gap"] = roles[0]
    store.upsert_user(user)
    return result


@app.post("/api/skills/gap")
def skill_gap(payload: SkillGapRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    context = profile_context(user)
    context["skills"] = payload.current_skills or context["skills"]
    context["resume_text"] = payload.resume_text or context["resume_text"]
    context["target_role"] = payload.target_career
    request_payload = {"context": context}
    schema = {
        "target_career": "",
        "job_readiness_percentage": 0,
        "current_skills": [],
        "required_skills": [],
        "missing_skills": [],
        "beginner_skills": [],
        "intermediate_skills": [],
        "advanced_skills": [],
        "priority_order": [{"skill": "", "priority": "", "reason": ""}],
        "recommended_learning_sequence": [],
        "project_suggestions": [],
        "interview_topics": [],
        "coding_practice_suggestions": [],
        "estimated_completion_timeline": "",
    }
    try:
        result = llm_cached(
            user,
            "skill_gap",
            request_payload,
            schema,
            f"Build an AI skill gap analysis. Required skills must be inferred from current industry expectations for {payload.target_career}. Profile JSON:\n{json.dumps(request_payload)}",
        )
    except HTTPException as exc:
        if not is_groq_auth_error(exc):
            raise
        result = local_skill_gap(context)
    user["latest_skill_gap"] = result
    store.upsert_user(user)
    return result


def question_embedding(text: str, size: int = 96) -> list[float]:
    vector = [0.0] * size
    for token in re.findall(r"[a-z0-9+#.]+", text.lower()):
        digest = hashlib.sha256(token.encode()).digest()
        vector[int.from_bytes(digest[:2], "big") % size] += -1.0 if digest[2] & 1 else 1.0
    length = sqrt(sum(value * value for value in vector)) or 1.0
    return [round(value / length, 6) for value in vector]


def question_similarity(text: str, embedding: list[float], previous: dict[str, Any]) -> float:
    old_text = str(previous.get("text") or "")
    lexical = SequenceMatcher(None, re.sub(r"\W+", " ", text.lower()), re.sub(r"\W+", " ", old_text.lower())).ratio()
    old_embedding = previous.get("embedding") or []
    cosine = sum(a * b for a, b in zip(embedding, old_embedding)) if old_embedding else 0.0
    return max(lexical, cosine)


def duplicate_question(candidate: dict[str, Any], history: list[dict[str, Any]]) -> bool:
    text = re.sub(r"\W+", " ", str(candidate.get("text") or "").lower()).strip()
    focus = str(candidate.get("focus_area") or "").lower().strip()
    embedding = candidate.get("embedding") or []
    for previous in history:
        old_text = re.sub(r"\W+", " ", str(previous.get("text") or "").lower()).strip()
        old_focus = str(previous.get("focus_area") or "").lower().strip()
        if text == old_text or (focus and focus == old_focus and question_similarity(text, embedding, previous) > 0.8):
            return True
    return False


def interview_topics(role: str) -> list[str]:
    catalog = store.courses_for_role(catalog_role(role))
    topics = [str(course.get("title") or "") for course in catalog]
    return clean_list([*topics, *role_required_skills(role)], 80) or [role, "technical fundamentals", "system design"]


def interview_rounds(session: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"id": index, "round_number": index, "title": title, "difficulty": difficulty, "count": QUESTIONS_PER_ROUND, "completed": sum(int(answer.get("round_number") or 0) == index for answer in session.get("answers", []))} for index, (title, difficulty) in enumerate(INTERVIEW_ROUNDS, 1)]


def interview_progress(session: dict[str, Any], round_number: int, question_number: int) -> dict[str, Any]:
    answered = len(session.get("answers", []))
    remaining = max(0, TOTAL_INTERVIEW_QUESTIONS - answered)
    round_answered = sum(int(item.get("round_number") or 0) == round_number for item in session.get("answers", []))
    return {"percent": round(answered / TOTAL_INTERVIEW_QUESTIONS * 100), "answered": answered, "total": TOTAL_INTERVIEW_QUESTIONS, "remaining": remaining, "round": min(round_number, len(INTERVIEW_ROUNDS)), "round_total": len(INTERVIEW_ROUNDS), "question": min(question_number, QUESTIONS_PER_ROUND), "round_question_total": QUESTIONS_PER_ROUND, "round_percent": round(round_answered / QUESTIONS_PER_ROUND * 100), "estimated_remaining_minutes": remaining * 2, "interview_seconds_remaining": remaining * 120, "round_seconds_remaining": max(0, (QUESTIONS_PER_ROUND - round_answered) * 120), "question_seconds_remaining": 120}


def interview_scoreboard(session: dict[str, Any]) -> dict[str, Any]:
    answers = session.get("answers", [])
    counts = {status: sum(item.get("status") == status for item in answers) for status in ("CORRECT", "WRONG", "SKIPPED", "TIMEOUT")}
    asked = len(answers)
    if sum(counts.values()) != asked:
        raise RuntimeError("Interview status invariant violated")
    scored = counts["CORRECT"] + counts["WRONG"]
    average_score = round(sum(int(item.get("score") or 0) for item in answers) / max(1, asked))
    return {"correct": counts["CORRECT"], "wrong": counts["WRONG"], "skipped": counts["SKIPPED"], "timeout": counts["TIMEOUT"], "answered": asked, "remaining": max(0, TOTAL_INTERVIEW_QUESTIONS - asked), "accuracy": round(counts["CORRECT"] / max(1, scored) * 100), "progress": round(asked / TOTAL_INTERVIEW_QUESTIONS * 100), "score": average_score}


def round_summary(session: dict[str, Any], round_number: int) -> dict[str, Any]:
    items = [item for item in session.get("answers", []) if int(item.get("round_number") or 0) == round_number]
    counts = {status: sum(item.get("status") == status for item in items) for status in ("CORRECT", "WRONG", "SKIPPED", "TIMEOUT")}
    criteria = lambda key: round(sum(int(item.get("criteria", {}).get(key) or 0) for item in items) / max(1, len(items)))
    return {"round": round_number, "title": INTERVIEW_ROUNDS[round_number - 1][0], "questions": len(items), "correct": counts["CORRECT"], "wrong": counts["WRONG"], "skipped": counts["SKIPPED"], "timeout": counts["TIMEOUT"], "accuracy": round(counts["CORRECT"] / max(1, counts["CORRECT"] + counts["WRONG"]) * 100), "average_time": round(sum(int(item.get("timeTaken") or 0) for item in items) / max(1, len(items))), "technical_score": criteria("Technical Accuracy"), "communication_score": criteria("Communication"), "confidence_score": criteria("Confidence")}


def previous_interview_questions(user: dict[str, Any], session: dict[str, Any]) -> list[dict[str, Any]]:
    previous: list[dict[str, Any]] = list(session.get("questions", []))
    for old_session in user.get("mock_interviews", [])[-10:]:
        if old_session.get("session_id") != session.get("session_id"):
            previous.extend(old_session.get("questions", []))
    return previous


def fallback_interview_question(session: dict[str, Any], round_number: int, question_number: int, topic: str, question_id: str) -> dict[str, Any]:
    title, difficulty = INTERVIEW_ROUNDS[round_number - 1]
    role, company = session["role"], session["company"]
    if round_number == 3:
        text = f"Using {session['language']}, design and implement solution {question_number} that applies {topic} to a production dataset of scale {question_number * 10000}. Explain complexity and edge cases."
        return {"id": question_id, "text": text, "problem_statement": text, "type": "coding", "starter_code": "", "constraints": ["Explain time and space complexity", "Handle empty and large inputs"], "test_cases": ["Typical input", "Empty input", "Large input"]}
    if round_number == 1:
        text = f"Give me a concise introduction for this {company} {role} interview, then connect one resume achievement to your experience with {topic}."
    elif round_number == 5:
        text = f"Choose a project from your resume that used {topic}. What did you personally own, what tradeoff did you make, and what measurable result followed?"
    elif round_number == 8:
        text = f"Tell me about a difficult collaboration while delivering {topic}-related work. How did you handle disagreement, and what would you change now?"
    elif round_number == 9:
        text = f"{company} rapid fire: state the most important production risk in {topic}, your first diagnostic step, and the tradeoff in your preferred fix."
    else:
        text = f"In a {company}-style {role} interview, how would you apply {topic} during {title.lower()}? Explain decisions, tradeoffs, risks, and a concrete example."
    return {"id": question_id, "text": text, "type": "text"}


def generate_interview_question(user: dict[str, Any], session: dict[str, Any], round_number: int, question_number: int) -> dict[str, Any]:
    title, difficulty = INTERVIEW_ROUNDS[round_number - 1]
    topics = interview_topics(session["role"])
    used = {str(question.get("focus_area", "")).split("/", 1)[0].strip().lower() for question in session.get("questions", []) if int(question.get("round_number") or 0) == round_number}
    available = [topic for topic in topics if topic.lower() not in used] or topics
    history = previous_interview_questions(user, session)
    for attempt in range(5):
        topic = secrets.choice(available)
        concept = f"{topic} / {title} / facet {question_number}"
        question_id = secrets.token_hex(8)
        schema = {"question": {"id": question_id, "text": "", "type": "coding" if round_number == 3 else "text", "difficulty": difficulty, "focus_area": topic, "timer_seconds": 120, "problem_statement": "", "starter_code": "", "constraints": [], "test_cases": [], "ideal_signals": []}}
        prior_text = [question.get("text", "") for question in history[-40:]]
        try:
            generated = groq_json(
                "You create realistic company-style interviews. Generate exactly one unique question, stay strictly within the target role, never repeat a prior concept, and return JSON only.",
                f"Company: {session['company']}. Target role: {session['role']}. Experience: {session['experience_level']}. Round {round_number}/9, question {question_number}/10: {title}. Difficulty: {difficulty}. Focus concept: {concept}. Coding language: {session['language']}. This is {'a coding problem with constraints and starter code' if round_number == 3 else 'a professional spoken interview question'}. Avoid all previous questions and concepts: {json.dumps(prior_text)}. Variation nonce: {secrets.token_hex(6)}.",
                schema,
                0.8,
            ).get("question", {})
        except HTTPException as exc:
            if not is_groq_auth_error(exc):
                raise
            generated = fallback_interview_question(session, round_number, question_number, topic, question_id)
        generated.update({"id": generated.get("id") or question_id, "round": title, "round_number": round_number, "question_number": question_number, "difficulty": difficulty, "focus_area": concept, "timer_seconds": max(30, int(generated.get("timer_seconds") or 120))})
        if round_number == 3:
            generated["type"] = "coding"
        else:
            generated["type"] = generated.get("type") if generated.get("type") in {"text", "mcq"} else "text"
        embedding = question_embedding(str(generated.get("text") or generated.get("problem_statement") or ""))
        generated["embedding"] = embedding
        if generated.get("text") and not duplicate_question(generated, history):
            return generated
    raise HTTPException(status_code=503, detail="Could not generate a unique interview question. Please start a new interview.")


def find_mock_interview(user: dict[str, Any], session_id: str) -> dict[str, Any]:
    session = next((item for item in user.get("mock_interviews", []) if item.get("session_id") == session_id), None)
    if not session:
        raise HTTPException(status_code=404, detail="Interview session not found. Start a new interview.")
    return session


def build_interview_report(user: dict[str, Any], session: dict[str, Any]) -> dict[str, Any]:
    answers = session.get("answers", [])
    scores = [int(item.get("overall") or 0) for item in answers]
    overall = round(sum(scores) / max(1, len(scores)))
    scoreboard = interview_scoreboard(session)
    skipped, timeout = scoreboard["skipped"], scoreboard["timeout"]
    correct, wrong = scoreboard["correct"], scoreboard["wrong"]
    rounds = {str(item.get("round")): int(item.get("overall") or 0) for item in answers}
    history = [{"session_id": old.get("session_id"), "role": old.get("role"), "date": old.get("created_at"), "score": old.get("overall", 0)} for old in user.get("mock_interviews", [])[-6:]]
    criterion_average = lambda key: round(sum(int(item.get("criteria", {}).get(key) or 0) for item in answers) / max(1, len(answers)))
    round_summaries = [round_summary(session, number) for number in range(1, len(INTERVIEW_ROUNDS) + 1) if any(int(item.get("round_number") or 0) == number for item in answers)]
    return {"overall": overall, "overall_questions": len(answers), "decision": "Interview Complete", "hiring_recommendation": "Strong Hire" if overall >= 80 else "Hire" if overall >= 65 else "Keep Practicing", "decision_reason": f"Completed {len(answers)} of {TOTAL_INTERVIEW_QUESTIONS} company-style questions for {session['role']} at {session['company']}.", "questions_correct": correct, "questions_wrong": wrong, "questions_skipped": skipped, "questions_timeout": timeout, "accuracy": scoreboard["accuracy"], "technical_score": criterion_average("Technical Accuracy"), "coding_score": round(sum(int(a.get("overall") or 0) for a in answers if a.get("round_number") == 3) / max(1, sum(a.get("round_number") == 3 for a in answers))), "hr_score": round(sum(int(a.get("overall") or 0) for a in answers if a.get("round_number") == 8) / max(1, sum(a.get("round_number") == 8 for a in answers))), "communication": criterion_average("Communication"), "confidence": criterion_average("Confidence"), "problem_solving": criterion_average("Problem Solving"), "behavior": round(sum(int(a.get("overall") or 0) for a in answers if a.get("round_number") == 8) / max(1, sum(a.get("round_number") == 8 for a in answers))), "role_readiness": overall, "rounds": rounds, "round_summaries": round_summaries, "items": answers, "strengths": clean_list([point for item in answers for point in item.get("strengths", [])], 200)[:8], "weaknesses": clean_list([point for item in answers for point in item.get("weaknesses", [])], 200)[:8], "recommended_learning_path": clean_list([point for item in answers for point in item.get("practice_plan", [])], 200)[:10], "areas_to_improve": clean_list([point for item in answers for point in item.get("missing_points", [])], 200)[:10], "skill_gap": {"missing_skills": [], "strong_skills": [], "topics_to_learn": []}, "weekly_plan": [], "company_readiness": [{"company": session["company"], "readiness": overall, "likely_level": session["experience_level"]}], "history": history}


@app.post("/api/interview/start")
def start_interview(payload: InterviewStartRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    session = {"session_id": secrets.token_hex(12), "role": payload.target_role.strip() or user.get("target_role") or "Software Engineer", "language": payload.language, "experience_level": payload.experience_level, "company": secrets.choice(INTERVIEW_COMPANIES), "created_at": now_iso(), "status": "In Progress", "questions": [], "answers": [], "current_round": 1, "current_question": 1}
    question = generate_interview_question(user, session, 1, 1)
    session["questions"].append(question)
    user.setdefault("mock_interviews", []).append(session)
    store.upsert_user(user)
    return {**session, "question": question, "progress": interview_progress(session, 1, 1), "scoreboard": interview_scoreboard(session), "rounds": interview_rounds(session)}


@app.post("/api/interview/answer")
def answer_interview(payload: InterviewAnswerRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    session = find_mock_interview(user, payload.session_id)
    question = next((item for item in session.get("questions", []) if item.get("id") == payload.question_id), None)
    if not question:
        raise HTTPException(status_code=404, detail="Interview question not found in this session.")
    if any(item.get("question_id") == payload.question_id for item in session.get("answers", [])):
        raise HTTPException(status_code=409, detail="This question has already been answered.")
    skipped = payload.action in {"skip", "timeout", "finish"}
    schema = {"overall": 0, "evaluation_confidence": 0, "correctness": "", "criteria": {"Technical Accuracy": 0, "Problem Solving": 0, "Communication": 0, "Confidence": 0, "Best Practices": 0, "Clarity": 0}, "strengths": [], "weaknesses": [], "missing_points": [], "mistakes": [], "suggestions": [], "ideal_answer": "", "industry_expected_answer": "", "reason": "", "practice_plan": []}
    if skipped:
        result = {**schema, "correctness": "Not Answered", "reason": "The question was skipped or timed out.", "ideal_answer": "Review the topic and practice a structured answer."}
    else:
        try:
            result = groq_json("You are a strict, fair company interviewer. Evaluate only the supplied answer and return JSON only.", f"Company: {session['company']}. Role: {session['role']}. Difficulty: {question['difficulty']}. Question: {question['text']}. Candidate answer: {payload.answer}. Candidate code: {payload.code}. Score correctness, technical depth, communication, confidence, problem solving, and best practices. Give strengths, weaknesses, ideal answer, and actionable improvement tips.", schema, 0.35)
        except HTTPException as exc:
            if not is_groq_auth_error(exc):
                raise
            result = local_interview_feedback(payload)
    if payload.action == "timeout":
        status = "TIMEOUT"
    elif payload.action in {"skip", "finish"}:
        status = "SKIPPED"
    else:
        confidence = int(result.get("evaluation_confidence") or result.get("overall") or 0)
        status = "CORRECT" if str(result.get("correctness", "")).lower() == "correct" and confidence >= 80 else "WRONG"
    result["correctness"] = "Correct" if status == "CORRECT" else "Not Answered" if status in {"SKIPPED", "TIMEOUT"} else "Incorrect"
    result.update({"interviewId": payload.session_id, "session_id": payload.session_id, "question_id": payload.question_id, "question": question["text"], "questionNumber": question["question_number"], "round": question["round"], "round_number": question["round_number"], "difficulty": question["difficulty"], "company": session["company"], "role": session["role"], "language": session["language"], "answer": payload.answer, "action": payload.action, "status": status, "correct": status == "CORRECT", "wrong": status == "WRONG", "skipped": status == "SKIPPED", "timeout": status == "TIMEOUT", "score": int(result.get("overall") or 0), "feedback": {"strengths": result.get("strengths", []), "weaknesses": result.get("weaknesses", []), "ideal_answer": result.get("ideal_answer", "")}, "timeTaken": payload.time_taken, "created_at": now_iso()})
    result["next_difficulty"] = INTERVIEW_ROUNDS[min(question["round_number"], 8)][1] if question["round_number"] < 9 else "Complete"
    session["answers"].append(result)
    finished = payload.action == "finish" or (question["round_number"] >= len(INTERVIEW_ROUNDS) and question["question_number"] >= QUESTIONS_PER_ROUND)
    next_question = None
    report = None
    completed_round = question["question_number"] >= QUESTIONS_PER_ROUND
    summary = round_summary(session, question["round_number"]) if completed_round else None
    if finished:
        session["status"] = "Completed"
        report = build_interview_report(user, session)
        session["overall"] = report["overall"]
        session["completed_at"] = now_iso()
    else:
        session["current_round"] = question["round_number"] + 1 if completed_round else question["round_number"]
        session["current_question"] = 1 if completed_round else question["question_number"] + 1
        next_question = generate_interview_question(user, session, session["current_round"], session["current_question"])
        session["questions"].append(next_question)
    store.upsert_user(user)
    return {**result, "next_question": next_question, "round_summary": summary, "report": report, "scoreboard": interview_scoreboard(session), "progress": interview_progress(session, session.get("current_round", 9), session.get("current_question", 10)), "session": {"status": session["status"], "rounds": interview_rounds(session)}}


@app.get("/api/interview/report")
def interview_report(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    sessions = user.get("mock_interviews", [])
    if not sessions:
        return {"overall": 0, "rounds": {}, "items": [], "recommendations": ["Start a mock interview and answer at least one question."]}
    return build_interview_report(user, sessions[-1])


@app.post("/api/github/analyze")
def analyze_github(
    payload: GitHubAnalysisRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    return github_analytics(payload.identifier, user, payload.target_role)


@app.get("/api/github/{username}")
def github_analytics(
    username: str,
    user: dict[str, Any] = Depends(current_user),
    target_role: str = "",
) -> dict[str, Any]:
    username = github_username(username)
    try:
        account_response = github_get(f"https://api.github.com/users/{username}")
        if account_response.status_code == 404:
            raise HTTPException(status_code=404, detail="GitHub user not found.")
        account_response.raise_for_status()
        repos_response = github_get(
            f"https://api.github.com/users/{username}/repos",
            params={"per_page": 100, "sort": "updated", "type": "owner"},
        )
        repos_response.raise_for_status()
        account = account_response.json()
        repos = repos_response.json()
    except HTTPException:
        raise
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Could not reach GitHub: {exc}") from exc
    context = profile_context(user)
    if target_role.strip():
        context["target_role"] = target_role.strip()
    result = github_portfolio_report(username, account, repos, context)
    user["github_username"] = username
    user["github_analysis"] = result
    store.upsert_user(user)
    return result


@app.get("/api/jobs")
def job_recommendations(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    context = profile_context(user)
    schema = {
        "target_role": "",
        "jobs": [
            {
                "title": "",
                "match_percentage": 0,
                "missing_skills": [],
                "required_skills": [],
                "salary_range": "",
                "recommended_improvements": [],
                "why_it_matches": "",
                "job_search_keywords": [],
            }
        ],
    }
    try:
        result = llm_cached(
            user,
            "jobs",
            context,
            schema,
            f"Recommend realistic job types/search targets from this profile. Do not invent company openings or apply URLs. Include match percentage, missing skills, required skills, salary range, and improvements. Profile JSON:\n{json.dumps(context)}",
        )
    except HTTPException as exc:
        if not is_groq_auth_error(exc):
            raise
        result = local_jobs(context)
    jobs = result.get("jobs", [])
    for job in jobs:
        job["match"] = job.get("match_percentage", job.get("match", 0))
        job["skills"] = job.get("required_skills", [])
        job["salary"] = job.get("salary_range", "")
    return {"jobs": jobs, "target_role": result.get("target_role") or user.get("target_role"), "total": len(jobs), "cached": result.get("cached", False)}


@app.post("/api/mentor")
def mentor(payload: MentorRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, str]:
    history = user.setdefault("mentor_messages", [])[-12:]
    messages = [
        {
            "role": "system",
            "content": (
                "You are CareerPilot AI Mentor, a friendly and experienced career mentor who speaks naturally, not like a chatbot. "
                "Answer only from the supplied career data and conversation: profile, target role, skills, resume ATS report, GitHub "
                "analytics, skill gaps, readiness, roadmap, interview results, learning progress, courses, and job recommendations. "
                "Never invent facts, scores, experience, projects, salaries, certifications, companies, or live openings. If needed data "
                "is absent, say exactly what is missing and ask at most one focused question. Detect the user's intent and tailor the "
                "answer to their target role. For profile requests cover identity, role, readiness, skills, gaps, progress, scores, and "
                "next goal. For learning, readiness, projects, resume, GitHub, or interview requests include the useful requested details "
                "from their data, but do not force a fixed template. Be supportive, specific, concise, and actionable. Never begin with "
                "'For your question', 'The best move is', or 'Based on available data'. Avoid repeating wording or paragraphs found in "
                "previous assistant messages. End with one brief, natural, encouraging sentence."
            ),
        },
        {"role": "user", "content": f"Saved career data JSON:\n{json.dumps(mentor_context(user), default=str)}"},
    ]
    for item in history:
        messages.append({"role": item["role"], "content": item["content"]})
    messages.append({"role": "user", "content": payload.message})
    try:
        reply = groq_completion(messages, temperature=0.35, max_tokens=1200)
    except HTTPException as exc:
        if not is_groq_auth_error(exc):
            raise
        reply = local_mentor_reply(payload.message, user)
    history.extend([{"role": "user", "content": payload.message}, {"role": "assistant", "content": reply}])
    user["mentor_messages"] = history[-14:]
    store.upsert_user(user)
    return {"reply": reply}
